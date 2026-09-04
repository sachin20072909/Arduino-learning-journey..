#!/usr/bin/env python3
"""Build the PhishGuard 16:9 hackathon pitch deck as an editable PPTX."""

from pathlib import Path

from lxml import etree
from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

BG = RGBColor(6, 9, 16)
CARD = RGBColor(16, 24, 40)
CARD_ALT = RGBColor(15, 23, 42)
LINE = RGBColor(30, 58, 74)
CYAN = RGBColor(34, 211, 238)
CYAN_SOFT = RGBColor(103, 232, 249)
TEXT = RGBColor(241, 245, 249)
MUTED = RGBColor(148, 163, 184)
DIM = RGBColor(100, 116, 139)
GREEN = RGBColor(52, 211, 153)
YELLOW = RGBColor(251, 191, 36)
RED = RGBColor(251, 113, 133)
BAD_LINE = RGBColor(80, 40, 52)
GOOD_LINE = RGBColor(24, 80, 64)

W = Inches(13.333)
H = Inches(7.5)
FONT = "Calibri"
ROOT = Path(__file__).resolve().parent
ASSETS = ROOT / "assets"
CACHE = ASSETS / "_fit"
CACHE.mkdir(parents=True, exist_ok=True)


def _set_run(run, size, color, bold=False, italic=False):
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.font.bold = bold
    run.font.italic = italic
    run.font.name = FONT
    rPr = run._r.get_or_add_rPr()
    for tag in ("a:latin", "a:cs", "a:ea"):
        el = rPr.find(qn(tag))
        if el is None:
            el = etree.SubElement(rPr, qn(tag))
        el.set("typeface", FONT)


def add_text(
    slide, l, t, w, h, text, size=14, color=TEXT, bold=False, italic=False,
    align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, margin=0.08,
):
    box = slide.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    tf.margin_left = Inches(margin)
    tf.margin_right = Inches(margin)
    tf.margin_top = Inches(0.04)
    tf.margin_bottom = Inches(0.02)
    try:
        tf._txBody.bodyPr.set(
            "anchor",
            {MSO_ANCHOR.TOP: "t", MSO_ANCHOR.MIDDLE: "ctr", MSO_ANCHOR.BOTTOM: "b"}[anchor],
        )
    except Exception:
        pass
    for i, line in enumerate(text.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_before = Pt(0)
        p.space_after = Pt(2)
        p.line_spacing = 1.08
        run = p.add_run()
        run.text = line
        _set_run(run, size, color, bold, italic)
    return box


def shape_fill_line(shape, fill, line=None, weight=1.0):
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    if line is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = line
        shape.line.width = Pt(weight)


def set_alpha(shape, pct):
    """pct: 0 transparent … 100 opaque."""
    solid = shape._element.spPr.find(qn("a:solidFill"))
    if solid is None:
        return
    srgb = solid.find(qn("a:srgbClr"))
    if srgb is None:
        return
    for child in list(srgb):
        if child.tag == qn("a:alpha"):
            srgb.remove(child)
    alpha = etree.SubElement(srgb, qn("a:alpha"))
    alpha.set("val", str(int(pct * 1000)))


def rect(slide, l, t, w, h, fill=CARD, line=LINE, radius=True, weight=1.15, alpha=None):
    kind = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    s = slide.shapes.add_shape(kind, Inches(l), Inches(t), Inches(w), Inches(h))
    shape_fill_line(s, fill, line, weight)
    if radius:
        try:
            s.adjustments[0] = 0.08
        except Exception:
            pass
    s.shadow.inherit = False
    if alpha is not None:
        set_alpha(s, alpha)
    return s


def oval(slide, l, t, w, h, fill, line=None):
    s = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(l), Inches(t), Inches(w), Inches(h))
    shape_fill_line(s, fill, line)
    return s


def fit_image(name, w_in, h_in, radius=28, dpi=160):
    src = ASSETS / name
    w, h = int(w_in * dpi), int(h_in * dpi)
    out = CACHE / f"{src.stem}_{w}x{h}_r{radius}.png"
    im = Image.open(src).convert("RGBA")
    tw, th = im.size
    scale = max(w / tw, h / th)
    nw, nh = int(tw * scale), int(th * scale)
    im = im.resize((nw, nh), Image.Resampling.LANCZOS)
    left, top = (nw - w) // 2, (nh - h) // 2
    im = im.crop((left, top, left + w, top + h))
    if radius > 0:
        mask = Image.new("L", (w, h), 0)
        ImageDraw.Draw(mask).rounded_rectangle((0, 0, w - 1, h - 1), radius=radius, fill=255)
        im.putalpha(mask)
    im.save(out, optimize=True)
    return str(out)


def pic(slide, name, l, t, w, h, radius=28):
    path = fit_image(name, w, h, radius=radius)
    return slide.shapes.add_picture(path, Inches(l), Inches(t), Inches(w), Inches(h))


def set_bg(slide):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = BG
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), W, H)
    shape_fill_line(bg, BG, None)
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), W, Inches(0.055))
    shape_fill_line(bar, CYAN, None)


def chrome(slide, n, total=12, light=False):
    color = TEXT if light else CYAN
    dim = RGBColor(203, 213, 225) if light else DIM
    add_text(slide, 0.48, 0.14, 6.8, 0.32, "PHISHGUARD  ·  ON-DEVICE ML", 11, color, True)
    add_text(slide, 11.2, 0.14, 1.7, 0.32, f"{n:02d}  /  {total:02d}", 11, dim, False, align=PP_ALIGN.RIGHT)


def kicker(slide, text, y=0.48):
    add_text(slide, 0.5, y, 12.2, 0.32, text.upper(), 12, CYAN, True)


def title(slide, text, y=0.78, h=1.15, size=32, w=12.3):
    add_text(slide, 0.5, y, w, h, text, size, TEXT, True)


def lede(slide, text, y, w=11.5):
    add_text(slide, 0.5, y, w, 0.55, text, 16, MUTED, False)


def pill(slide, l, t, w, h, text, fg=TEXT, bg=CARD_ALT, ln=LINE):
    rect(slide, l, t, w, h, bg, ln, True, 1.0)
    add_text(slide, l, t + 0.04, w, h - 0.04, text, 12, fg, True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, margin=0.04)


def node(slide, l, t, w, h, text, sub="", accent=False, fg=None):
    ln = CYAN if accent else LINE
    fill = RGBColor(10, 32, 42) if accent else CARD
    rect(slide, l, t, w, h, fill, ln, True, 1.25 if accent else 1.0)
    color = fg or (CYAN_SOFT if accent else TEXT)
    if sub:
        add_text(slide, l, t + 0.08, w, 0.28, text, 12, color, True, align=PP_ALIGN.CENTER, margin=0.04)
        add_text(slide, l, t + 0.32, w, 0.28, sub, 10, DIM, False, align=PP_ALIGN.CENTER, margin=0.04)
    else:
        add_text(slide, l, t, w, h, text, 13, color, True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, margin=0.04)


def arrow(slide, l, t, text="→"):
    add_text(slide, l, t, 0.32, 0.4, text, 18, CYAN, True, align=PP_ALIGN.CENTER)


def card_block(slide, l, t, w, h, heading, body, ln=LINE):
    rect(slide, l, t, w, h, CARD, ln)
    add_text(slide, l + 0.06, t + 0.12, w - 0.16, 0.36, heading, 15, TEXT, True)
    add_text(slide, l + 0.06, t + 0.48, w - 0.16, h - 0.58, body, 13, MUTED, False)


def bullets(slide, l, t, w, h, items, size=14, color=MUTED):
    box = slide.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.04)
    tf.margin_right = Inches(0.06)
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_before = Pt(3)
        p.space_after = Pt(4)
        p.line_spacing = 1.08
        run = p.add_run()
        run.text = "▸  " + item
        _set_run(run, size, color, False)
    return box


def notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text


def build():
    prs = Presentation()
    prs.slide_width = W
    prs.slide_height = H
    blank = prs.slide_layouts[6]

    # ========== 1 TITLE ==========
    s = prs.slides.add_slide(blank)
    set_bg(s)
    pic(s, "hero_browser_shield.png", 0, 0, 13.333, 7.5, radius=0)
    rect(s, 0, 0, 8.15, 7.5, BG, None, False, alpha=78)
    rect(s, 0, 0, 13.333, 0.055, CYAN, None, False)
    chrome(s, 1)
    kicker(s, "Hackathon pitch  ·  Privacy-first web security")
    add_text(s, 0.5, 0.95, 7.3, 1.9, "Browser-Based\nPhishing Detector", 36, TEXT, True)
    add_text(
        s, 0.5, 2.95, 7.1, 0.95,
        "On-device machine learning for real-time web security.\nDetect phishing websites without sending user data to the cloud.",
        16, MUTED, False,
    )
    pill(s, 0.5, 4.05, 1.7, 0.38, "On-device ML", CYAN_SOFT, RGBColor(10, 32, 42), CYAN)
    pill(s, 2.32, 4.05, 1.45, 0.38, "Real-time", TEXT, CARD, LINE)
    pill(s, 3.88, 4.05, 1.7, 0.38, "Privacy-first", TEXT, CARD, LINE)
    pill(s, 5.68, 4.05, 1.7, 0.38, "Lightweight", TEXT, CARD, LINE)
    add_text(s, 0.5, 4.7, 2.3, 0.22, "TEAM", 10, DIM, True)
    add_text(s, 0.5, 4.92, 2.3, 0.35, "[Team Name]", 14, TEXT, False)
    add_text(s, 2.9, 4.7, 4.4, 0.22, "MEMBERS", 10, DIM, True)
    add_text(s, 2.9, 4.92, 4.4, 0.35, "[Member 1]  ·  [Member 2]  ·  [Member 3]", 14, TEXT, False)
    add_text(s, 0.5, 5.4, 7.0, 0.22, "EVENT", 10, DIM, True)
    add_text(s, 0.5, 5.62, 7.0, 0.35, "[Hackathon Name]", 14, TEXT, False)
    notes(s, "Open on the privacy hook. Introduce team. ~30s.")

    # ========== 2 PROBLEM ==========
    s = prs.slides.add_slide(blank)
    set_bg(s)
    chrome(s, 2)
    kicker(s, "The problem")
    title(s, "One Click Can Expose Everything.", y=0.78, h=0.55, size=30, w=7.4)
    pic(s, "problem_fake_login.png", 7.55, 0.85, 5.3, 4.35, radius=22)
    cards = [
        ("Convincing fakes", "Phishing sites imitate banks, email, and cloud logins people already trust."),
        ("High-value theft", "One form submit can leak passwords, cards, and personal details."),
        ("Cloud trade-offs", "Many detectors send URLs or page data to remote servers."),
        ("No time to wait", "A warning must appear before the user hits “Sign in”."),
    ]
    for i, (hd, bd) in enumerate(cards):
        x, y = 0.5 + (i % 2) * 3.45, 1.5 + (i // 2) * 1.75
        card_block(s, x, y, 3.3, 1.62, hd, bd)
    node(s, 7.55, 5.3, 1.2, 0.42, "User", fg=TEXT)
    add_text(s, 8.75, 5.3, 0.28, 0.42, "→", 16, CYAN, True, anchor=MSO_ANCHOR.MIDDLE)
    node(s, 9.05, 5.3, 1.7, 0.42, "Fake site", fg=TEXT)
    add_text(s, 10.75, 5.3, 0.28, 0.42, "→", 16, RED, True, anchor=MSO_ANCHOR.MIDDLE)
    node(s, 11.05, 5.3, 1.8, 0.42, "Stolen", fg=RGBColor(254, 205, 211), accent=True)
    rect(s, 0.5, 5.9, 12.35, 0.95, RGBColor(10, 32, 42), CYAN, True, 1.3)
    add_text(
        s, 0.65, 5.98, 12.05, 0.8,
        "Key question   ·   How can we detect phishing in real time while keeping user data on the device?",
        16, CYAN_SOFT, True, anchor=MSO_ANCHOR.MIDDLE,
    )
    notes(s, "Make the privacy tension clear. End on the key question. ~35s.")

    # ========== 3 SOLUTION ==========
    s = prs.slides.add_slide(blank)
    set_bg(s)
    chrome(s, 3)
    kicker(s, "Our solution")
    title(s, "Phishing Detection, Directly Inside the Browser.", y=0.72, h=0.7, size=28, w=7.5)
    add_text(s, 0.5, 1.42, 7.4, 0.55, "A Chrome extension that scores the current page locally — then shows a clear, human verdict.", 15, MUTED)
    pic(s, "solution_extension.png", 8.05, 0.7, 4.8, 2.85, radius=20)
    nodes = [
        (0.5, "Website", "Current tab", False),
        (3.05, "Feature extraction", "URL + page signals", False),
        (5.6, "On-device ML", "No cloud inference", True),
        (8.15, "Risk score", "0–100%", False),
        (10.7, "Verdict", "3 simple states", False),
    ]
    for l, a, b, acc in nodes:
        node(s, l, 3.65, 2.35, 0.78, a, b, acc)
    for l in (2.78, 5.33, 7.88, 10.43):
        arrow(s, l, 3.82)
    pill(s, 0.5, 4.55, 1.55, 0.36, "●  Safe", GREEN, CARD, LINE)
    pill(s, 2.2, 4.55, 2.15, 0.36, "●  Suspicious", YELLOW, CARD, LINE)
    pill(s, 4.5, 4.55, 2.45, 0.36, "●  Likely Phishing", RED, CARD, LINE)
    feats = [
        ("Local", "Runs on the user’s device."),
        ("Private", "No page data required in the cloud."),
        ("Real-time", "Scores the page as it loads."),
        ("Lightweight", "Compact model for the browser."),
        ("Transparent", "Shows the signals behind the call."),
    ]
    for i, (hd, bd) in enumerate(feats):
        card_block(s, 0.5 + i * 2.5, 5.08, 2.38, 1.95, hd, bd)
    notes(s, "Product slide. Hit on-device, private, real-time, lightweight. ~40s.")

    # ========== 4 ARCHITECTURE ==========
    s = prs.slides.add_slide(blank)
    set_bg(s)
    chrome(s, 4)
    kicker(s, "Architecture")
    title(s, "From URL to Verdict in Milliseconds", y=0.72, h=0.5, size=28)
    lede(s, "Every step — extraction, scoring, and the final warning — happens inside the browser. Inference does not require a cloud round-trip.", 1.22)
    rect(s, 0.5, 1.85, 12.35, 5.15, CARD, LINE)
    node(s, 1.1, 2.1, 3.3, 0.7, "Chrome browser")
    arrow(s, 4.5, 2.23)
    node(s, 4.85, 2.1, 3.3, 0.7, "Browser extension")
    arrow(s, 8.25, 2.23)
    node(s, 8.6, 2.1, 3.55, 0.7, "Feature extraction", accent=True)
    node(s, 1.3, 3.05, 5.15, 0.95, "URL features", "Length, host, TLD, lookalikes, IP, special chars")
    node(s, 6.9, 3.05, 5.15, 0.95, "Page content features", "Forms, HTTPS, iframes, brand mismatch")
    node(s, 0.9, 4.25, 3.7, 0.7, "30–40 engineered features")
    arrow(s, 4.7, 4.37)
    node(s, 5.05, 4.25, 3.5, 0.7, "On-device ML model", accent=True)
    arrow(s, 8.65, 4.37)
    node(s, 9.0, 4.25, 3.2, 0.7, "Phishing probability")
    node(s, 1.1, 5.25, 3.5, 0.7, "Safe", fg=GREEN)
    node(s, 4.95, 5.25, 3.5, 0.7, "Suspicious", fg=YELLOW)
    node(s, 8.8, 5.25, 3.5, 0.7, "Likely phishing", fg=RED)
    add_text(s, 0.7, 6.15, 12.0, 0.5, "ALL INFERENCE RUNS LOCALLY  ·  NO CLOUD ROUND-TRIP REQUIRED", 12, DIM, True, align=PP_ALIGN.CENTER)
    notes(s, "Walk the pipeline once. Emphasize local inference. ~40s.")

    # ========== 5 SIGNALS ==========
    s = prs.slides.add_slide(blank)
    set_bg(s)
    chrome(s, 5)
    kicker(s, "Signals")
    title(s, "30–40 Signals Behind Every Decision", y=0.72, h=0.5, size=28)
    pic(s, "signals_scan.png", 0.5, 1.28, 12.35, 1.55, radius=18)
    rect(s, 0.5, 3.0, 6.05, 4.0, CARD, LINE)
    add_text(s, 0.7, 3.12, 5.7, 0.22, "URL-BASED FEATURES", 11, CYAN, True)
    add_text(s, 0.7, 3.36, 5.7, 0.32, "What the address reveals", 16, TEXT, True)
    bullets(s, 0.7, 3.72, 5.6, 3.1, [
        "URL length and nested path structure",
        "Subdomains and suspicious TLDs",
        "Special characters and encoding tricks",
        "IP address used instead of a domain",
        "Typosquatting / lookalike brands",
        "Suspicious URL structures overall",
    ], 14, MUTED)
    rect(s, 6.8, 3.0, 6.05, 4.0, CARD, LINE)
    add_text(s, 7.0, 3.12, 5.7, 0.22, "PAGE CONTENT FEATURES", 11, CYAN, True)
    add_text(s, 7.0, 3.36, 5.7, 0.32, "What the page is asking for", 16, TEXT, True)
    bullets(s, 7.0, 3.72, 5.6, 3.1, [
        "Login forms and password fields",
        "Credential-harvesting indicators",
        "HTTP / non-secure connections",
        "Embedded iframes",
        "Brand vs. content mismatch",
        "Suspicious form destinations",
    ], 14, MUTED)
    notes(s, "Pick 2 URL + 2 page examples a judge can picture. ~35s.")

    # ========== 6 ML ==========
    s = prs.slides.add_slide(blank)
    set_bg(s)
    chrome(s, 6)
    kicker(s, "Machine learning  ·  model-selection stage")
    title(s, "Small Model. Fast Decision.", y=0.72, h=0.5, size=30, w=7.6)
    add_text(s, 0.5, 1.25, 7.5, 0.7, "Trained offline, then shipped with the extension. Inference stays on-device. Final architecture is still being selected.", 15, MUTED)
    pic(s, "ml_compact.png", 8.15, 0.7, 4.7, 2.55, radius=20)
    rect(s, 0.5, 2.05, 7.4, 2.35, CARD, LINE)
    add_text(s, 0.7, 2.15, 7.0, 0.22, "CANDIDATE A", 11, CYAN, True)
    add_text(s, 0.7, 2.4, 7.0, 0.32, "Gradient-boosted decision tree", 17, TEXT, True)
    bullets(s, 0.7, 2.78, 7.0, 1.45, [
        "Strong on tabular URL / page features",
        "Stays compact after training",
        "Easier to inspect which signals matter",
    ], 14)
    rect(s, 8.15, 3.4, 4.7, 2.0, CARD, LINE)
    add_text(s, 8.35, 3.5, 4.35, 0.22, "CANDIDATE B", 11, CYAN, True)
    add_text(s, 8.35, 3.75, 4.35, 0.32, "Compact neural network", 16, TEXT, True)
    bullets(s, 8.35, 4.12, 4.35, 1.15, [
        "Small enough for browser runtimes",
        "Fast CPU forward-pass / ONNX · TF.js",
    ], 13)
    node(s, 0.5, 4.6, 2.35, 0.62, "30–40 features")
    arrow(s, 2.88, 4.7)
    node(s, 3.2, 4.6, 2.7, 0.62, "Selected compact model", accent=True)
    arrow(s, 5.95, 4.7)
    node(s, 6.28, 4.6, 1.7, 0.62, "Score")
    pill(s, 0.5, 5.45, 1.85, 0.38, "Trained offline", TEXT, CARD, LINE)
    pill(s, 2.5, 5.45, 2.2, 0.38, "No cloud inference", TEXT, CARD, LINE)
    pill(s, 4.85, 5.45, 2.15, 0.38, "Low-latency design", TEXT, CARD, LINE)
    pill(s, 7.15, 5.45, 2.7, 0.38, "Feature-based = explainable", TEXT, CARD, LINE)
    add_text(s, 0.5, 6.0, 12.3, 0.55, "Model choice is still open — we are comparing these two compact options, not claiming a final architecture.", 14, DIM)
    notes(s, "Be honest: model not locked. Sell compactness and locality. ~35s.")

    # ========== 7 PRIVACY ==========
    s = prs.slides.add_slide(blank)
    set_bg(s)
    chrome(s, 7)
    kicker(s, "Privacy by design")
    title(s, "Your Data Stays on Your Device.", y=0.72, h=0.5, size=28)
    pic(s, "privacy_ondevice.png", 0.5, 1.28, 12.35, 1.7, radius=18)
    rect(s, 0.5, 3.15, 6.05, 3.85, CARD, BAD_LINE, True, 1.4)
    add_text(s, 0.7, 3.28, 5.7, 0.22, "TRADITIONAL CLOUD DETECTION", 11, RED, True)
    add_text(s, 0.7, 3.52, 5.7, 0.35, "Page data leaves the browser", 16, TEXT, True)
    for i, lab in enumerate(["User", "Website information", "Cloud server", "Detection", "Result"]):
        node(s, 0.75, 3.95 + i * 0.55, 5.55, 0.48, lab, fg=MUTED)
    rect(s, 6.8, 3.15, 6.05, 3.85, CARD, GOOD_LINE, True, 1.4)
    add_text(s, 7.0, 3.28, 5.7, 0.22, "OUR APPROACH", 11, GREEN, True)
    add_text(s, 7.0, 3.52, 5.7, 0.35, "Inference never needs the cloud", 16, TEXT, True)
    for i, lab in enumerate(["User", "Website information", "Local ML model", "Result"]):
        acc = lab == "Local ML model"
        node(s, 7.05, 3.95 + i * 0.52, 5.55, 0.46, lab, accent=acc, fg=CYAN_SOFT if acc else TEXT)
    add_text(s, 7.05, 6.15, 5.55, 0.7, "No URL or page content uploaded  ·  Works offline-ish  ·  Faster local decisions", 13, MUTED)
    notes(s, "This is the differentiator. Contrast cloud vs local. ~40s.")

    # ========== 8 UX ==========
    s = prs.slides.add_slide(blank)
    set_bg(s)
    chrome(s, 8)
    kicker(s, "User experience  ·  example UI  ·  not a benchmark")
    title(s, "What the User Sees", y=0.72, h=0.45, size=30)
    add_text(s, 0.5, 1.2, 7.5, 0.35, "Simple language. One glance. Built for a non-technical user.", 15, MUTED)
    rect(s, 0.5, 1.65, 5.5, 5.3, RGBColor(16, 24, 38), CYAN, True, 1.4)
    add_text(s, 0.7, 1.8, 3.4, 0.3, "🛡  PhishGuard", 14, TEXT, True)
    pill(s, 4.15, 1.82, 1.6, 0.32, "LIVE EXAMPLE", RED, RGBColor(40, 18, 28), RED)
    add_text(s, 0.7, 2.22, 5.1, 0.28, "http://paypa1-secure-login.xyz", 12, MUTED, False)
    add_text(s, 0.7, 2.55, 5.1, 0.4, "⚠  LIKELY PHISHING", 20, RED, True)
    add_text(s, 0.7, 3.05, 2.2, 0.25, "RISK SCORE", 11, DIM, True)
    add_text(s, 3.7, 2.95, 2.0, 0.5, "94%", 28, RED, True, align=PP_ALIGN.RIGHT)
    rect(s, 0.7, 3.5, 5.05, 0.14, RGBColor(30, 41, 59), None, True)
    rect(s, 0.7, 3.5, 4.75, 0.14, RGBColor(251, 113, 133), None, True)
    add_text(s, 0.7, 3.8, 5.1, 0.25, "WHY THIS PAGE", 11, DIM, True)
    bullets(s, 0.7, 4.1, 5.1, 2.4, [
        "Suspicious URL structure",
        "Brand similarity detected",
        "Credential form found",
        "Non-secure connection",
    ], 14, MUTED)
    statuses = [
        (GREEN, "SAFE — Low Risk", "Looks consistent with a legitimate site."),
        (YELLOW, "SUSPICIOUS — Review Carefully", "Some signals don’t add up. Pause before you sign in."),
        (RED, "LIKELY PHISHING — Avoid This Site", "Do not enter passwords or payment details."),
    ]
    for i, (col, hd, bd) in enumerate(statuses):
        y = 1.65 + i * 1.75
        rect(s, 6.25, y, 6.6, 1.6, CARD, LINE)
        oval(s, 6.45, y + 0.68, 0.22, 0.22, col)
        add_text(s, 6.85, y + 0.22, 5.75, 0.4, hd, 16, col, True)
        add_text(s, 6.85, y + 0.68, 5.75, 0.7, bd, 14, MUTED, False)
    notes(s, "94% is an example verdict for the popup, not a measured accuracy. ~35s.")

    # ========== 9 EVALUATION ==========
    s = prs.slides.add_slide(blank)
    set_bg(s)
    chrome(s, 9)
    kicker(s, "Evaluation  ·  targets, not claimed results")
    title(s, "We Measure What Actually Matters.", y=0.72, h=0.5, size=28)
    metrics = [
        ("PRECISION", "Target", "Of flagged sites, how many were truly phishing?"),
        ("RECALL", "Priority", "Of real phishing sites, how many did we catch?"),
        ("F1 SCORE", "Balance", "Single view of precision and recall together."),
        ("LATENCY", "ms-class", "Time from page load to on-device verdict."),
        ("MODEL SIZE", "Compact", "Small enough to ship inside the extension."),
    ]
    for i, (k, n, b) in enumerate(metrics):
        x = 0.5 + i * 2.5
        rect(s, x, 1.45, 2.38, 2.55, CARD, LINE)
        add_text(s, x + 0.12, 1.6, 2.14, 0.25, k, 11, DIM, True)
        add_text(s, x + 0.12, 1.95, 2.14, 0.45, n, 20, CYAN_SOFT, True)
        add_text(s, x + 0.12, 2.5, 2.14, 1.3, b, 13, MUTED, False)
    rect(s, 0.5, 4.2, 12.35, 1.25, RGBColor(10, 32, 42), CYAN, True, 1.2)
    add_text(
        s, 0.7, 4.35, 12.0, 1.0,
        "Recall is the safety metric.  A missed phishing page (false negative) can expose a password.\nWe will tune thresholds to catch attacks first, then reduce false alarms.",
        16, TEXT, False, anchor=MSO_ANCHOR.MIDDLE,
    )
    pill(s, 0.5, 5.7, 3.15, 0.4, "Target detection quality: 95%+", CYAN_SOFT, RGBColor(10, 32, 42), CYAN)
    pill(s, 3.8, 5.7, 3.05, 0.4, "Not yet experimentally verified", MUTED, CARD, LINE)
    pill(s, 7.0, 5.7, 2.55, 0.4, "Held-out test set planned", MUTED, CARD, LINE)
    notes(s, "Do not invent numbers. Say target 95%+ clearly. Stress recall. ~30s.")

    # ========== 10 DATA ==========
    s = prs.slides.add_slide(blank)
    set_bg(s)
    chrome(s, 10)
    kicker(s, "Data  ·  planned public sources")
    title(s, "Built and Tested Against Real-World Data", y=0.72, h=0.5, size=28, w=8.0)
    pic(s, "dataset_world.png", 8.2, 0.7, 4.65, 2.15, radius=18)
    sources = [
        ("PHISHING", "PhishTank", "Community-verified phishing URLs as malicious examples."),
        ("PHISHING", "OpenPhish", "Additional live phishing feeds for coverage and freshness."),
        ("BENIGN", "Tranco", "Popular legitimate websites as the safe class."),
    ]
    for i, (k, hd, bd) in enumerate(sources):
        x = 0.5 + i * 4.2
        rect(s, x, 3.0, 4.0, 1.55, CARD, LINE)
        add_text(s, x + 0.18, 3.1, 3.65, 0.22, k, 11, CYAN, True)
        add_text(s, x + 0.18, 3.35, 3.65, 0.32, hd, 18, TEXT, True)
        add_text(s, x + 0.18, 3.7, 3.65, 0.7, bd, 13, MUTED, False)
    pipe = [(0.5, "Dataset"), (3.05, "Feature extraction"), (5.6, "Train / test split"), (8.15, "Model training"), (10.7, "Held-out eval")]
    for i, (x, lab) in enumerate(pipe):
        node(s, x, 4.7, 2.35, 0.62, lab, accent=(lab == "Held-out eval"))
        if i < 4:
            arrow(s, x + 2.28, 4.8)
    card_block(s, 0.5, 5.5, 6.05, 1.45, "Consistent features", "The same 30–40 signals are computed for every URL and page snapshot.")
    card_block(s, 6.8, 5.5, 6.05, 1.45, "Generalization first", "We score the held-out set — not training accuracy — so results reflect new sites.")
    notes(s, "Name the three sources. Stress held-out evaluation. ~30s.")

    # ========== 11 LIMITS ==========
    s = prs.slides.add_slide(blank)
    set_bg(s)
    chrome(s, 11)
    kicker(s, "Honesty & roadmap")
    title(s, "Honest About the Gaps. Ready to Improve.", y=0.72, h=0.55, size=28)
    rect(s, 0.5, 1.45, 6.05, 4.2, CARD, LINE)
    add_text(s, 0.7, 1.6, 5.7, 0.3, "CURRENT LIMITATIONS", 11, YELLOW, True)
    bullets(s, 0.7, 2.0, 5.6, 3.4, [
        "Cannot reliably check live domain age without an external service",
        "Focused on web-based phishing, not every scam channel",
        "Sophisticated attacks may still evade feature-based detection",
        "Dynamic pages can change after the first analysis",
        "False positives are possible",
    ], 14, MUTED)
    rect(s, 6.8, 1.45, 6.05, 4.2, CARD, LINE)
    add_text(s, 7.0, 1.6, 5.7, 0.3, "FUTURE IMPROVEMENTS", 11, CYAN, True)
    bullets(s, 7.0, 2.0, 5.6, 3.4, [
        "Deeper behavioral analysis after load",
        "Stronger typosquatting detection",
        "Optional reputation lists (user-controlled)",
        "Continuous model improvement",
        "Richer social-engineering patterns and more browsers",
    ], 14, MUTED)
    rect(s, 0.5, 5.85, 12.35, 0.95, RGBColor(10, 32, 42), CYAN, True, 1.2)
    add_text(
        s, 0.65, 5.95, 12.05, 0.75,
        "Hackathon-ready core: extract signals locally → score with a compact model → warn before submit. Extras stay optional so privacy remains the default.",
        15, CYAN_SOFT, False, anchor=MSO_ANCHOR.MIDDLE,
    )
    notes(s, "Name 2 limits, 2 next steps, then the feasible core. ~30s.")

    # ========== 12 CLOSE ==========
    s = prs.slides.add_slide(blank)
    set_bg(s)
    pic(s, "impact_close.png", 0, 0, 13.333, 7.5, radius=0)
    rect(s, 0, 0, 13.333, 7.5, BG, None, False, alpha=62)
    rect(s, 0, 0, 13.333, 0.055, CYAN, None, False)
    chrome(s, 12)
    kicker(s, "Impact")
    add_text(s, 0.5, 0.85, 12.3, 0.7, "Security Without Sacrificing Privacy.", 32, TEXT, True, align=PP_ALIGN.CENTER)
    trio = [
        ("DETECT", "See it in time", "Identify suspicious phishing patterns in real time, in the tab the user is already on."),
        ("PROTECT", "Warn before submit", "Give a clear Safe / Suspicious / Phishing call before sensitive data is entered."),
        ("PRESERVE", "Keep it local", "Run detection and ML inference on the device — not on someone else’s server."),
    ]
    for i, (k, hd, bd) in enumerate(trio):
        x = 0.5 + i * 4.2
        rect(s, x, 1.7, 4.0, 2.15, CARD, LINE, alpha=88)
        add_text(s, x + 0.2, 1.85, 3.6, 0.25, k, 12, CYAN, True)
        add_text(s, x + 0.2, 2.15, 3.6, 0.4, hd, 18, TEXT, True)
        add_text(s, x + 0.2, 2.6, 3.6, 1.0, bd, 14, MUTED, False)
    rect(s, 1.4, 4.1, 10.5, 0.7, RGBColor(10, 32, 42), CYAN, True, 1.3, alpha=90)
    add_text(
        s, 1.55, 4.15, 10.2, 0.6,
        "Effective phishing detection doesn't have to mean sending user data to the cloud.",
        16, CYAN_SOFT, True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE,
    )
    add_text(s, 0.5, 5.05, 12.3, 0.55, "Thank You", 32, TEXT, True, align=PP_ALIGN.CENTER)
    add_text(s, 0.5, 5.55, 12.3, 0.35, "Questions?", 18, CYAN, True, align=PP_ALIGN.CENTER)
    add_text(
        s, 0.5, 6.05, 12.3, 0.4,
        "[Team Name]   ·   github.com/[your-team]/phishguard   ·   [contact@email]",
        14, MUTED, False, align=PP_ALIGN.CENTER,
    )
    notes(s, "Close on Detect / Protect / Preserve. Invite questions. ~25s.")

    out = ROOT / "PhishGuard-Hackathon-Pitch.pptx"
    prs.save(str(out))
    print(out)


if __name__ == "__main__":
    build()
