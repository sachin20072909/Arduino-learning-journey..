"""
build_deck.py
=============

Builds ``WiFiSense_Hackathon_Deck.pptx`` — a 10-slide, 16:9 dark-tech
hackathon pitch deck for the WiFiSense project.

Run ``prepare_assets.py`` first (it produces the crops in assets/processed),
then:

    python3 deck/build_deck.py

The deck is intentionally text-light and diagram-heavy: every slide carries
one idea, one hero visual, and at most a handful of short phrases.
"""

from __future__ import annotations

import math
import random
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.oxml import parse_xml
from pptx.oxml.ns import nsdecls
from pptx.util import Emu, Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
ASSETS = ROOT / "assets" / "processed"
OUT = ROOT / "WiFiSense_Hackathon_Deck.pptx"

# ------------------------------------------------------------------ theme

BG        = "05070C"   # slide background
PANEL     = "0B1220"   # card fill
PANEL_2   = "0E1728"   # raised card
LINE      = "1B2740"   # hairline borders
CYAN      = "22D3EE"   # primary accent
CYAN_D    = "0EA5E9"   # secondary accent
VIOLET    = "8B5CF6"   # tertiary accent
GREEN     = "34D399"
AMBER     = "FBBF24"
RED       = "FB7185"
TEXT      = "EAF2FF"
MUTED     = "9AAECB"
DIM       = "647890"

FONT = "Segoe UI"

SW, SH = 13.333, 7.5          # slide size (inches)
M = 0.75                      # side margin
CW = SW - 2 * M               # content width = 11.833


# ------------------------------------------------------------- primitives


def rgb(value: str) -> RGBColor:
    return RGBColor.from_string(value)


def _alpha_el(parent, alpha: float):
    """Append <a:alpha val="..."/> (thousandths of a percent) to a colour el."""
    from pptx.oxml.ns import qn as _qn
    a = parent.makeelement(_qn("a:alpha"), {})
    a.set("val", str(int(round(alpha * 100000))))
    parent.append(a)
    return a


def solid(shape, color: str, alpha: float | None = None):
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(color)
    if alpha is not None:
        el = shape.fill._xPr.find(qn("a:srgbClr"))
        if el is not None:
            _alpha_el(el, alpha)
    return shape


def outline(shape, color: str, width: float = 1.0, alpha: float | None = None):
    shape.line.color.rgb = rgb(color)
    shape.line.width = Pt(width)
    if alpha is not None:
        ln = shape.line._get_or_add_ln()
        el = ln.find(qn("a:srgbClr"))
        if el is not None:
            _alpha_el(el, alpha)
    return shape


def no_line(shape):
    shape.line.fill.background()
    return shape


def no_fill(shape):
    shape.fill.background()
    return shape


def glow(shape, color: str, radius_pt: float = 10, alpha: float = 0.55):
    """Add an <a:glow/> effect (PowerPoint's soft neon halo)."""
    spPr = shape._element.spPr
    effect_lst = spPr.find(qn("a:effectLst"))
    if effect_lst is None:
        effect_lst = spPr.makeelement(qn("a:effectLst"), {})
        spPr.append(effect_lst)
    g = effect_lst.makeelement(qn("a:glow"), {})
    g.set("rad", str(int(radius_pt * 12700)))
    clr = g.makeelement(qn("a:srgbClr"), {"val": color})
    a = clr.makeelement(qn("a:alpha"), {})
    a.set("val", str(int(alpha * 100000)))
    clr.append(a)
    g.append(clr)
    effect_lst.append(g)
    return shape


def rect(slide, x, y, w, h, fill: str | None = PANEL, fill_alpha: float | None = None,
         line: str | None = None, lw: float = 1.0, line_alpha: float | None = None,
         radius: float | None = None, shape=MSO_SHAPE.ROUNDED_RECTANGLE):
    if radius is None:
        shape = MSO_SHAPE.RECTANGLE
    shp = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if fill is None:
        no_fill(shp)
    else:
        solid(shp, fill, fill_alpha)
    if line is None:
        no_line(shp)
    else:
        outline(shp, line, lw, line_alpha)
    if radius is not None:
        try:
            shp.adjustments[0] = min(0.5, radius / (min(w, h) / 2))
        except (IndexError, KeyError, TypeError):
            pass
    shp.shadow.inherit = False
    shp.text_frame.text = ""
    return shp


def oval(slide, x, y, w, h, fill=None, fill_alpha=None, line=None, lw=1.0,
         line_alpha=None):
    return rect(slide, x, y, w, h, fill, fill_alpha, line, lw, line_alpha,
                shape=MSO_SHAPE.OVAL)


def txt(slide, x, y, w, h, text, size=14, color=TEXT, bold=False,
        font=FONT, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
        spacing=None, line_spacing=None, italic=False, caps=False,
        alpha=None):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    p = tf.paragraphs[0]
    p.alignment = align
    if line_spacing:
        p.line_spacing = line_spacing
    run = p.add_run()
    run.text = text.upper() if caps else text
    f = run.font
    f.name, f.size, f.bold, f.italic = font, Pt(size), bold, italic
    f.color.rgb = rgb(color)
    if alpha is not None:
        rPr = run._r.get_or_add_rPr()
        el = rPr.find(qn("a:solidFill")).find(qn("a:srgbClr"))
        _alpha_el(el, alpha)
    if spacing:
        rPr = run._r.get_or_add_rPr()
        rPr.set("spc", str(int(spacing * 100)))
    return box


def picture(slide, name: str, x, y, w=None, h=None):
    path = ASSETS / name
    if not path.exists():
        raise FileNotFoundError(path)
    return slide.shapes.add_picture(str(path), Inches(x), Inches(y),
                                    Inches(w) if w else None,
                                    Inches(h) if h else None)


def poly(slide, pts, color=CYAN, width=1.6, alpha=None):
    """Draw an open polyline (used for the CSI waveform)."""
    fb = slide.shapes.build_freeform(Inches(pts[0][0]), Inches(pts[0][1]))
    fb.add_line_segments([(Inches(px), Inches(py)) for px, py in pts[1:]])
    shp = fb.convert_to_shape()
    no_fill(shp)
    outline(shp, color, width, alpha)
    return shp


def arrow_right(slide, x, y, w, h, color=CYAN, alpha=None):
    shp = rect(slide, x, y, w, h, fill=color, fill_alpha=alpha,
               shape=MSO_SHAPE.RIGHT_ARROW)
    return shp


def arrow_down(slide, x, y, w, h, color=CYAN, alpha=None):
    return rect(slide, x, y, w, h, fill=color, fill_alpha=alpha,
                shape=MSO_SHAPE.DOWN_ARROW)


# ------------------------------------------------------------ slide frame


def blank(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg = slide.background
    bg.fill.solid()
    bg.fill.fore_color.rgb = rgb(BG)
    return slide


def backdrop(slide, name: str, alpha: float = 1.0):
    """Full-bleed background picture."""
    pic = picture(slide, name, 0, 0, SW, SH)
    if alpha < 1.0:
        # dim by stacking a translucent scrim
        scrim = rect(slide, 0, 0, SW, SH, fill=BG, fill_alpha=1 - alpha)
        scrim.shadow.inherit = False
    return pic


def header(slide, eyebrow: str, title: str, kicker: str | None = None):
    txt(slide, M, 0.40, 8.0, 0.24, eyebrow, size=10.5, color=CYAN, bold=True,
        spacing=2.2)
    # tiny accent square before the eyebrow
    rect(slide, M - 0.22, 0.455, 0.10, 0.10, fill=CYAN)
    txt(slide, M, 0.70, CW, 0.55, title, size=30, color=TEXT, bold=True)
    y = 1.32
    if kicker:
        txt(slide, M, y, CW - 2.6, 0.30, kicker, size=12.5, color=MUTED)
        y += 0.42
    line = rect(slide, M, y - 0.10, CW, 0.012, fill=LINE)
    return y + 0.12


def footer(slide, number: int, total: int = 10, label: str = "WiFiSense"):
    txt(slide, M, SH - 0.46, 6.0, 0.22,
        f"{label}  ·  Privacy-Preserving Human Activity Detection Using WiFi CSI",
        size=8, color=DIM, spacing=0.6)
    txt(slide, SW - M - 1.2, SH - 0.46, 1.2, 0.22,
        f"{number:02d} / {total:02d}", size=9, color=CYAN, bold=True,
        align=PP_ALIGN.RIGHT, spacing=1.2)


def chip(slide, x, y, w, h, label, color=CYAN, size=11, bold=False,
         fill_alpha=0.10, line_alpha=0.38, text_color=None, spacing=None):
    rect(slide, x, y, w, h, fill=color, fill_alpha=fill_alpha,
         line=color, lw=0.9, line_alpha=line_alpha, radius=min(0.16, h / 2))
    txt(slide, x + 0.10, y, w - 0.20, h, label, size=size,
        color=text_color or color, bold=bold, align=PP_ALIGN.CENTER,
        anchor=MSO_ANCHOR.MIDDLE, spacing=spacing)
    return (x + w, y)


def card(slide, x, y, w, h, fill=PANEL, line=LINE, radius=0.14,
         fill_alpha=None, line_alpha=None, accent: str | None = None,
         accent_side="left"):
    shp = rect(slide, x, y, w, h, fill=fill, fill_alpha=fill_alpha,
               line=line, lw=1.0, line_alpha=line_alpha, radius=radius)
    if accent:
        if accent_side == "left":
            rect(slide, x, y + radius * 0.6, 0.045, h - radius * 1.2, fill=accent)
        else:
            rect(slide, x, y, w, 0.045, fill=accent)
    return shp


def bullets(slide, x, y, w, items, gap=0.16, dot_color=CYAN, size=11.5,
            title_size=13.5, row_h=None):
    """items: list of (title, subtitle)."""
    cy = y
    for title, sub in items:
        h = row_h or (0.62 if sub else 0.34)
        oval(slide, x, cy + 0.10, 0.09, 0.09, fill=dot_color)
        txt(slide, x + 0.26, cy - 0.02, w - 0.26, 0.26, title, size=title_size,
            color=TEXT, bold=True)
        if sub:
            txt(slide, x + 0.26, cy + 0.26, w - 0.26, 0.36, sub, size=size,
                color=MUTED, line_spacing=1.05)
        cy += h + gap
    return cy


# ================================================================= slides


def slide_title(prs):
    s = blank(prs)
    backdrop(s, "s01_hero.png")

    # vertical light accent on the left edge
    rect(s, 0, 0, 0.055, SH, fill=CYAN)

    txt(s, M, 1.42, 7.5, 0.26, "HACKATHON PROJECT  ·  ESP32  ·  WIFI CSI SENSING",
        size=11, color=CYAN, bold=True, spacing=3.0)
    txt(s, M, 1.70, 8.6, 1.28, "WiFiSense", size=76, color=TEXT, bold=True,
        spacing=-1.2)
    txt(s, M, 3.02, 9.9, 0.40,
        "Privacy-Preserving Human Activity Detection Using WiFi CSI",
        size=19.5, color=CYAN, bold=False, spacing=0.2)

    # tagline
    tag_w = 4.30
    rect(s, M, 3.72, tag_w, 0.52, fill=CYAN, fill_alpha=0.10, line=CYAN,
         lw=1.1, line_alpha=0.45, radius=0.26)
    txt(s, M, 3.72, tag_w, 0.52, "Sense movement.  Not identities.", size=15,
        color=CYAN, bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE,
        spacing=1.0)

    txt(s, M, 4.52, 7.6, 0.60,
        "Two or three ESP32 boards read the WiFi channel instead of a camera.\n"
        "A lightweight model learns what human motion looks like — and what it does not.",
        size=13, color=MUTED, line_spacing=1.25)

    chips = ["2–3 × ESP32 boards", "No camera · No mic · No wearables",
             "On-device ML inference"]
    cx = M
    for label in chips:
        w = 0.115 * len(label) + 0.52
        chip(s, cx, 5.42, w, 0.40, label, color=CYAN, size=10.5, spacing=0.5)
        cx += w + 0.20

    # bottom rule + team line
    rect(s, M, 6.42, CW, 0.012, fill=LINE, line_alpha=0.6)
    txt(s, M, 6.58, 6.0, 0.26, "Team WiFiSense", size=11, color=TEXT, bold=True)
    txt(s, M + 1.45, 6.60, 6.0, 0.26,
        "ESP-IDF  ·  Python  ·  NumPy / Pandas  ·  scikit-learn  ·  FastAPI",
        size=10.5, color=DIM, spacing=0.6)
    txt(s, SW - M - 3.2, 6.58, 3.2, 0.26, "Live demo · 90 seconds",
        size=10.5, color=CYAN, bold=True, align=PP_ALIGN.RIGHT, spacing=1.0)
    return s


def slide_problem(prs):
    s = blank(prs)
    y = header(s, "01  —  THE PROBLEM", "Cameras are not always the answer",
               "Vision-based monitoring breaks down exactly where people care most about privacy.")

    # left: comparison artwork
    ix, iy, iw = M, y + 0.05, 6.10
    picture(s, "s02_problem.png", ix, iy, iw, iw / 2.0)
    rect(s, ix, iy, iw, iw / 2.0, fill=None, line=LINE, lw=1.0, radius=0.14)
    txt(s, ix, iy + iw / 2.0 + 0.16, iw, 0.24,
        "Cameras watch people.  WiFi sensing watches the air.",
        size=11, color=MUTED, align=PP_ALIGN.CENTER, spacing=0.4)

    # right: pain points
    rx = M + 6.55
    rw = CW - 6.55
    items = [
        ("Privacy", "Continuous video of people is intrusive — and hard to justify at home."),
        ("Blind spots", "A camera only sees where it points. Corners, shelves and furniture hide movement."),
        ("Darkness & glare", "Night, smoke, backlight and occlusion all degrade vision-based sensing."),
        ("Wearables don't stick", "Bands get forgotten, go uncharged, or are simply refused."),
    ]
    cy = y + 0.08
    for i, (t, sub) in enumerate(items):
        h = 0.86
        card(s, rx, cy, rw, h, radius=0.12,
             accent=[RED, AMBER, CYAN, VIOLET][i])
        oval(s, rx + 0.30, cy + 0.30, 0.26, 0.26, fill=[RED, AMBER, CYAN, VIOLET][i],
             fill_alpha=0.18, line=[RED, AMBER, CYAN, VIOLET][i], lw=0.9,
             line_alpha=0.5)
        txt(s, rx + 0.30, cy + 0.30, 0.26, 0.26, str(i + 1), size=11,
            color=[RED, AMBER, CYAN, VIOLET][i], bold=True, align=PP_ALIGN.CENTER,
            anchor=MSO_ANCHOR.MIDDLE)
        txt(s, rx + 0.72, cy + 0.16, rw - 0.95, 0.26, t, size=14.5, color=TEXT,
            bold=True)
        txt(s, rx + 0.72, cy + 0.44, rw - 0.95, 0.36, sub, size=11, color=MUTED,
            line_spacing=1.1)
        cy += h + 0.15

    # bottom strip
    strip_y = 6.30
    rect(s, M, strip_y, CW, 0.62, fill=CYAN, fill_alpha=0.08, line=CYAN,
         lw=1.0, line_alpha=0.30, radius=0.10)
    rect(s, M, strip_y, 0.05, 0.62, fill=CYAN)
    txt(s, M + 0.34, strip_y, CW - 0.6, 0.62,
        "We still need to know whether a space is occupied — without recording who is in it.",
        size=14.5, color=TEXT, bold=True, anchor=MSO_ANCHOR.MIDDLE, spacing=0.3)
    footer(s, 2)
    return s


def slide_solution(prs):
    s = blank(prs)
    y = header(s, "02  —  OUR SOLUTION", "Don't watch the room.  Read the WiFi.",
               "Instead of watching people with cameras, we observe how human movement changes WiFi signals.")

    # signal chain
    nodes = ["ESP32 TX", "WiFi Signal", "Human Body", "ESP32 RX",
             "CSI Stream", "ML Model", "Activity State"]
    n = len(nodes)
    gap = 0.30
    nw = (CW - gap * (n - 1)) / n
    cy = y + 0.10
    nh = 0.80
    for i, label in enumerate(nodes):
        nx = M + i * (nw + gap)
        is_last = i == n - 1
        is_ml = i == 5
        col = CYAN if is_last else (VIOLET if is_ml else CYAN_D)
        rect(s, nx, cy, nw, nh, fill=col, fill_alpha=0.10 if is_last else 0.06,
             line=col, lw=1.1, line_alpha=0.55 if is_last else 0.35, radius=0.10)
        if i in (0, 3):
            rect(s, nx + nw / 2 - 0.16, cy + 0.10, 0.32, 0.14, fill=col,
                 fill_alpha=0.75)
        txt(s, nx + 0.06, cy, nw - 0.12, nh, label, size=12.5, color=TEXT,
            bold=is_last, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE,
            spacing=0.3)
        if i < n - 1:
            arrow_right(s, nx + nw + 0.055, cy + nh / 2 - 0.055, gap - 0.11,
                        0.11, color=CYAN, alpha=0.65)

    # hero band
    by = cy + nh + 0.34
    bh = 2.42
    picture(s, "s03_band.png", M, by, CW, bh)
    rect(s, M, by, CW, bh, fill=None, line=LINE, lw=1.0, radius=0.14)
    txt(s, M + 0.30, by + 0.22, 6.2, 0.30,
        "THE ROOM IS THE SENSOR", size=12, color=CYAN, bold=True, spacing=2.0)
    txt(s, M + 0.30, by + 0.56, 6.6, 0.86,
        "A body moving through the link absorbs and scatters\nradio energy — the receiver sees it as CSI.",
        size=13, color=MUTED, line_spacing=1.2)

    # three bottom facts
    fy = by + bh + 0.26
    fw = (CW - 2 * 0.26) / 3
    facts = [
        (CYAN, "Low cost", "Two ESP32 boards, USB cables and a laptop."),
        (GREEN, "Privacy by design", "No images, no audio, no identity — only numbers."),
        (VIOLET, "Beyond line of sight", "Radio waves bend around furniture and work in the dark."),
    ]
    for i, (col, t, sub) in enumerate(facts):
        fx = M + i * (fw + 0.26)
        card(s, fx, fy, fw, 0.86, radius=0.10, accent=col)
        txt(s, fx + 0.28, fy + 0.12, fw - 0.5, 0.24, t, size=13, color=col,
            bold=True)
        txt(s, fx + 0.28, fy + 0.38, fw - 0.5, 0.36, sub, size=10.5,
            color=MUTED, line_spacing=1.1)
    footer(s, 3)
    return s


def slide_csi(prs):
    s = blank(prs)
    y = header(s, "03  —  HOW IT WORKS", "WiFi CSI, step by step",
               "CSI = Channel State Information — how the signal arrives, not the data it carries.")

    # left: waveform artwork
    iw, ih = 5.30, 3.62
    picture(s, "s04_csi.png", M, y + 0.06, iw, ih)
    rect(s, M, y + 0.06, iw, ih, fill=None, line=LINE, lw=1.0, radius=0.14)
    txt(s, M, y + 0.06 + ih + 0.16, iw, 0.52,
        "Each line is one subcarrier.  Human motion leaves a time-structured\n"
        "fingerprint; a quiet room leaves a flat one.",
        size=11, color=MUTED, line_spacing=1.15)

    # right: five steps
    rx = M + 5.85
    rw = CW - 5.85
    steps = [
        ("Packets on air", "An ESP32 transmits ordinary WiFi frames across the room."),
        ("Signal travels", "Energy reaches the receiver directly and via reflections off walls and objects."),
        ("A body disturbs it", "A moving person absorbs, reflects and scatters part of that energy."),
        ("Receiver logs CSI", "Per packet: amplitude and phase for each OFDM subcarrier — ~50 numbers."),
        ("ML reads the pattern", "Models look at how those numbers evolve over seconds, not at one reading."),
    ]
    cy = y + 0.06
    for i, (t, sub) in enumerate(steps):
        h = 0.70
        card(s, rx, cy, rw, h, radius=0.10)
        oval(s, rx + 0.26, cy + 0.19, 0.32, 0.32, fill=CYAN, fill_alpha=0.14,
             line=CYAN, lw=1.0, line_alpha=0.5)
        txt(s, rx + 0.26, cy + 0.19, 0.32, 0.32, str(i + 1), size=13,
            color=CYAN, bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        txt(s, rx + 0.74, cy + 0.11, rw - 1.0, 0.24, t, size=14, color=TEXT,
            bold=True)
        txt(s, rx + 0.74, cy + 0.37, rw - 1.0, 0.32, sub, size=10.5,
            color=MUTED, line_spacing=1.08)
        cy += h + 0.13
        if i < len(steps) - 1:
            arrow_down(s, rx + 0.35, cy - 0.09, 0.14, 0.10, color=CYAN,
                       alpha=0.45)
    footer(s, 4)
    return s


def slide_false_positive(prs):
    s = blank(prs)
    y = header(s, "04  —  THE HARD PROBLEM",
               "A changed signal is not always a human",
               "Bags, chairs, doors and fans also disturb the channel. Treating 'signal moved' as 'person' is the classic mistake.")

    icons = [("icon_person.png", "Person walking", CYAN),
             ("icon_bag.png", "Bag placed", AMBER),
             ("icon_chair.png", "Chair moved", AMBER),
             ("icon_fan.png", "Fan spinning", AMBER),
             ("icon_door.png", "Door opened", AMBER)]
    n = len(icons)
    gap = 0.22
    cw = (CW - gap * (n - 1)) / n
    ch = 2.52
    cy = y + 0.08
    for i, (img, label, col) in enumerate(icons):
        cx = M + i * (cw + gap)
        card(s, cx, cy, cw, ch, radius=0.12,
             accent=col if i == 0 else None)
        icon_size = 1.16
        picture(s, img, cx + (cw - icon_size) / 2, cy + 0.20, icon_size,
                icon_size)
        txt(s, cx + 0.12, cy + 1.46, cw - 0.24, 0.26, label, size=13,
            color=TEXT if i == 0 else MUTED, bold=True, align=PP_ALIGN.CENTER)
        # equals sign
        txt(s, cx + 0.12, cy + 1.78, cw - 0.24, 0.24, "=", size=15, color=DIM,
            bold=True, align=PP_ALIGN.CENTER)
        chip(s, cx + 0.28, cy + 2.04, cw - 0.56, 0.32, "CSI changes",
             color=col, size=10, fill_alpha=0.12, line_alpha=0.35, spacing=0.4)

    # callout
    ky = cy + ch + 0.24
    rect(s, M, ky, CW, 0.72, fill=CYAN, fill_alpha=0.09, line=CYAN, lw=1.1,
         line_alpha=0.40, radius=0.12)
    rect(s, M, ky, 0.055, 0.72, fill=CYAN)
    txt(s, M + 0.38, ky, CW - 0.7, 0.72,
        "CSI detects changes in the radio environment — not humans directly.",
        size=19, color=TEXT, bold=True, anchor=MSO_ANCHOR.MIDDLE, spacing=0.2)

    # three defences
    fy = ky + 0.90
    fw = (CW - 2 * 0.24) / 3
    defences = [
        (CYAN, "Temporal patterns", "Seconds of behaviour, never a single reading."),
        (VIOLET, "Balanced training data", "Human events AND non-human events in the same dataset."),
        (GREEN, "Multiple links", "2–3 ESP32 links cross-check a change from different angles."),
    ]
    for i, (col, t, sub) in enumerate(defences):
        fx = M + i * (fw + 0.24)
        card(s, fx, fy, fw, 0.80, radius=0.10, accent=col)
        txt(s, fx + 0.26, fy + 0.10, fw - 0.46, 0.24, t, size=13, color=col,
            bold=True)
        txt(s, fx + 0.26, fy + 0.35, fw - 0.46, 0.36, sub, size=10.5,
            color=MUTED, line_spacing=1.1)
    footer(s, 5)
    return s


def slide_pipeline(prs):
    s = blank(prs)
    y = header(s, "05  —  AI / ML PIPELINE",
               "From raw signal to a confidence score",
               "A small, explainable pipeline that runs on a laptop today and on the board tomorrow.")

    stages = [
        ("Raw CSI data", "Amplitude + phase, per subcarrier, per packet", CYAN_D, False),
        ("Noise filtering", "Drop empty packets, spikes and idle bursts", CYAN_D, False),
        ("Preprocessing", "Phase sanitised, normalised, windowed into segments", CYAN_D, False),
        ("Feature extraction", "Variance, Doppler energy, cross-link correlation", CYAN_D, False),
        ("ML classification", "Random Forest / SVM / tiny CNN — lightweight on purpose", VIOLET, False),
        ("Human · Object · Stable", "A three-way decision, never a hard yes / no", VIOLET, False),
        ("Live dashboard + alert", "Status, activity and a confidence score", CYAN, True),
    ]
    px, pw = M, 4.30
    cy = y + 0.05
    sh = 0.46
    arrow_gap = 0.24
    for i, (t, sub, col, hot) in enumerate(stages):
        card(s, px, cy, pw, sh, radius=0.08,
             fill=col if hot else PANEL, fill_alpha=0.14 if hot else None,
             line=col, line_alpha=0.55 if hot else 0.28, accent=col)
        txt(s, px + 0.30, cy, pw - 0.45, sh, t, size=13,
            color=TEXT if not hot else col, bold=hot,
            anchor=MSO_ANCHOR.MIDDLE, spacing=0.3)
        txt(s, px + pw + 0.12, cy - 0.02, 1.56, sh + 0.04, sub, size=8.5,
            color=DIM, anchor=MSO_ANCHOR.MIDDLE, line_spacing=1.05)
        if i < len(stages) - 1:
            arrow_down(s, px + pw / 2 - 0.075, cy + sh + 0.03, 0.15, 0.11,
                       color=CYAN, alpha=0.55)
        cy += sh + arrow_gap

    # right: ML artwork + probability readout
    rx = M + 6.10
    rw = CW - 6.10
    picture(s, "s06_ml.png", rx, y + 0.05, rw, 3.10)
    rect(s, rx, y + 0.05, rw, 3.10, fill=None, line=LINE, lw=1.0, radius=0.14)
    txt(s, rx + 0.30, y + 0.28, rw - 0.6, 0.26, "PATTERN, NOT PIXELS",
        size=11.5, color=CYAN, bold=True, spacing=2.2)

    py = y + 3.34
    card(s, rx, py, rw, 1.60, radius=0.12, accent=VIOLET)
    txt(s, rx + 0.32, py + 0.14, rw - 0.6, 0.26,
        "Output is a probability, not a verdict", size=14, color=TEXT, bold=True)
    probs = [("Human", 0.87, CYAN), ("Object", 0.11, AMBER),
             ("Stable", 0.02, GREEN)]
    by = py + 0.48
    for label, val, col in probs:
        txt(s, rx + 0.32, by, 1.0, 0.20, label, size=10.5, color=MUTED)
        bar_x = rx + 1.28
        bar_w = rw - 2.35
        rect(s, bar_x, by + 0.045, bar_w, 0.11, fill=LINE)
        rect(s, bar_x, by + 0.045, max(0.04, bar_w * val), 0.11, fill=col)
        txt(s, rx + rw - 0.72, by, 0.50, 0.20, f"{val:.2f}", size=10.5,
            color=col, bold=True, align=PP_ALIGN.RIGHT)
        by += 0.29
    txt(s, rx + 0.32, py + 1.30, rw - 0.6, 0.26,
        "Below threshold the system says “uncertain” — it never guesses.",
        size=9.5, color=DIM)
    footer(s, 6)
    return s


def slide_architecture(prs):
    s = blank(prs)
    backdrop(s, "s07_arch_bg.png")
    y = header(s, "06  —  SYSTEM ARCHITECTURE",
               "From radio hardware to a live dashboard",
               "Two to three ESP32 links feed one laptop; raw CSI never leaves the room.")

    col_w = 4.10
    mid_w = 1.30
    right_w = CW - col_w - mid_w - 0.60
    lx = M
    mx = lx + col_w + 0.30
    rx = mx + mid_w + 0.30

    row_h = 0.92
    row_gap = 0.30

    left_rows = [
        ("ESP32  Transmitter", "Broadcasts standard WiFi frames", "box_esp32.png", CYAN),
        ("WiFi Signal  ·  2.4 GHz", "Travels through walls, air and people", None, CYAN_D),
        ("Human  /  Environment", "Absorbs, reflects, scatters the energy", "box_rf.png", AMBER),
        ("ESP32  Receiver(s)  × 2–3", "Reports CSI per packet, per subcarrier", "box_esp32.png", CYAN),
    ]
    right_rows = [
        ("Laptop  ·  Edge Processing", "Parses, filters and windows the CSI stream", None, CYAN_D),
        ("ML Model", "Random Forest / SVM / TinyML — scores each window", "box_ml.png", VIOLET),
        ("Web Dashboard", "FastAPI backend + lightweight web UI", None, CYAN),
        ("Alerts  ·  Status  ·  Logs", "Room state, activity and confidence, live", None, GREEN),
    ]

    cy = y + 0.06
    for i, (t, sub, img, col) in enumerate(left_rows):
        card(s, lx, cy, col_w, row_h, radius=0.10, fill=PANEL,
             fill_alpha=0.72, accent=col)
        if img:
            picture(s, img, lx + 0.012, cy + 0.012, col_w - 0.024,
                    row_h - 0.024)
            rect(s, lx + 0.012, cy + 0.012, col_w - 0.024, row_h - 0.024,
                 fill=BG, fill_alpha=0.55)
        txt(s, lx + 0.32, cy + 0.16, col_w - 0.6, 0.28, t, size=14,
            color=TEXT, bold=True)
        txt(s, lx + 0.32, cy + 0.46, col_w - 0.6, 0.34, sub, size=10.5,
            color=MUTED, line_spacing=1.08)
        if i < len(left_rows) - 1:
            arrow_down(s, lx + col_w / 2 - 0.08, cy + row_h + 0.05, 0.16, 0.12,
                       color=CYAN, alpha=0.5)
        cy += row_h + row_gap

    cy = y + 0.06
    for i, (t, sub, img, col) in enumerate(right_rows):
        card(s, rx, cy, right_w, row_h, radius=0.10, fill=PANEL,
             fill_alpha=0.72, accent=col)
        if img:
            picture(s, img, rx + 0.012, cy + 0.012, right_w - 0.024,
                    row_h - 0.024)
            rect(s, rx + 0.012, cy + 0.012, right_w - 0.024, row_h - 0.024,
                 fill=BG, fill_alpha=0.55)
        txt(s, rx + 0.32, cy + 0.16, right_w - 0.6, 0.28, t, size=14,
            color=TEXT, bold=True)
        txt(s, rx + 0.32, cy + 0.46, right_w - 0.6, 0.34, sub, size=10.5,
            color=MUTED, line_spacing=1.08)
        if i < len(right_rows) - 1:
            arrow_down(s, rx + right_w / 2 - 0.08, cy + row_h + 0.05, 0.16, 0.12,
                       color=VIOLET, alpha=0.5)
        cy += row_h + row_gap

    # middle: CSI stream arrow from sensing column to processing column
    mid_cy = y + 0.06 + (row_h * 4 + row_gap * 3) / 2
    rect(s, mx, mid_cy - 0.28, mid_w, 0.56, fill=CYAN, fill_alpha=0.10,
         line=CYAN, lw=1.0, line_alpha=0.35, radius=0.08)
    arrow_right(s, mx + 0.10, mid_cy - 0.20, mid_w - 0.20, 0.40, color=CYAN,
                alpha=0.75)
    txt(s, mx - 0.02, mid_cy + 0.38, mid_w + 0.04, 0.62,
        "CSI stream → laptop:  USB / UART / Wi-Fi.  Numbers only, never images.",
        size=8, color=MUTED, align=PP_ALIGN.CENTER, line_spacing=1.12)

    # privacy badge
    py = 6.34
    rect(s, M, py, CW, 0.54, fill=GREEN, fill_alpha=0.09, line=GREEN, lw=1.0,
         line_alpha=0.35, radius=0.10)
    txt(s, M + 0.34, py, CW - 0.7, 0.54,
        "No camera  ·  No microphone  ·  No video or audio is ever captured, stored or transmitted — only numbers.",
        size=13, color=GREEN, bold=True, anchor=MSO_ANCHOR.MIDDLE, spacing=0.3)
    footer(s, 7)
    return s


def slide_usecases(prs):
    s = blank(prs)
    y = header(s, "07  —  WHERE IT FITS", "Real-world use cases",
               "Places where a camera is unwanted, inappropriate, or simply not enough.")

    cases = [
        ("tile_privacy.png", "Privacy-sensitive rooms", "Bedrooms, bathrooms, clinics — presence without pictures.", CYAN),
        ("tile_elderly.png", "Elderly & fall-risk care", "Flag long stillness or a sudden drop. No cameras in the home.", VIOLET),
        ("tile_building.png", "Smart buildings", "Occupancy-aware lighting, HVAC and real energy savings.", CYAN_D),
        ("tile_restricted.png", "Restricted areas", "After-hours movement in labs, stores or stock rooms.", AMBER),
        ("tile_night.png", "Night-time detection", "Works in complete darkness — no light, no infrared.", VIOLET),
        ("tile_occupancy.png", "Occupancy monitoring", "Desks and meeting rooms, counted in real time.", GREEN),
    ]
    gap = 0.26
    cw = (CW - gap * 2) / 3
    ch = 1.96
    img_h = 1.10
    row_gap = 0.18
    for i, (img, t, sub, col) in enumerate(cases):
        r, c = divmod(i, 3)
        cx = M + c * (cw + gap)
        cy = y + 0.06 + r * (ch + row_gap)
        card(s, cx, cy, cw, ch, radius=0.12, accent=col)
        picture(s, img, cx + 0.012, cy + 0.012, cw - 0.024, img_h)
        rect(s, cx + 0.012, cy + 0.012, cw - 0.024, img_h, fill=BG,
             fill_alpha=0.12)
        txt(s, cx + 0.26, cy + img_h + 0.14, cw - 0.5, 0.24, t, size=14,
            color=TEXT, bold=True)
        txt(s, cx + 0.26, cy + img_h + 0.40, cw - 0.5, 0.44, sub, size=10.5,
            color=MUTED, line_spacing=1.12)

    # honesty strip
    sy = y + 0.06 + 2 * (ch + row_gap) + 0.10
    rect(s, M, sy, CW, 0.56, fill=AMBER, fill_alpha=0.08, line=AMBER, lw=1.0,
         line_alpha=0.28, radius=0.10)
    txt(s, M + 0.34, sy, CW - 0.7, 0.56,
        "Honest scope: results depend on walls, layout and furniture — we do not claim reliable "
        "sensing in every environment.",
        size=12, color=AMBER, anchor=MSO_ANCHOR.MIDDLE, spacing=0.2)
    footer(s, 8)
    return s


def slide_demo(prs):
    s = blank(prs)
    y = header(s, "08  —  LIVE DEMO", "Five steps, ninety seconds",
               "Two ESP32 boards on a table, a laptop, and one honest failure test.")

    steps = [
        ("1", "Empty room, boards running", "No human activity", GREEN, "stable"),
        ("2", "A person walks in and stands", "Human detected", CYAN, "present"),
        ("3", "The person walks across the link", "Walking detected", CYAN, "moving"),
        ("4", "The person leaves the room", "Room stable", GREEN, "stable"),
        ("5", "A bag / chair is moved instead", "Not classified as human", AMBER, "rejected"),
    ]
    lx, lw = M, 6.42
    cy = y + 0.06
    row_h = 0.72
    gap = 0.14
    for num, action, expect, col, _ in steps:
        card(s, lx, cy, lw, row_h, radius=0.10, accent=col)
        oval(s, lx + 0.26, cy + 0.19, 0.30, 0.30, fill=col, fill_alpha=0.16,
             line=col, lw=1.0, line_alpha=0.55)
        txt(s, lx + 0.26, cy + 0.19, 0.30, 0.30, num, size=12.5, color=col,
            bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        txt(s, lx + 0.72, cy + 0.13, lw - 3.0, 0.26, action, size=13.5,
            color=TEXT, bold=True)
        txt(s, lx + 0.72, cy + 0.41, lw - 3.0, 0.24, "expected output",
            size=9.5, color=DIM, italic=True)
        w_expect = 1.95
        chip(s, lx + lw - w_expect - 0.22, cy + 0.19, w_expect, 0.32, expect,
             color=col, size=10, fill_alpha=0.14, line_alpha=0.40, spacing=0.3)
        cy += row_h + gap

    # ---- live dashboard mock (built natively so text stays crisp) ----
    rx = M + 6.78
    rw = CW - 6.78
    panel_h = 4.80
    py = y + 0.06
    card(s, rx, py, rw, panel_h, radius=0.14, fill=PANEL, fill_alpha=0.9,
         line=CYAN, line_alpha=0.30)
    # header bar
    rect(s, rx + 0.012, py + 0.012, rw - 0.024, 0.50, fill=CYAN,
         fill_alpha=0.08)
    oval(s, rx + 0.26, py + 0.20, 0.11, 0.11, fill=CYAN)
    txt(s, rx + 0.46, py + 0.13, 3.0, 0.28, "WiFiSense  ·  Live Monitor",
        size=11.5, color=TEXT, bold=True, spacing=0.4)
    chip(s, rx + rw - 1.72, py + 0.17, 1.50, 0.28, "●  LIVE", color=GREEN,
         size=9.5, fill_alpha=0.12, line_alpha=0.35, spacing=0.8)

    # status pill
    sy = py + 0.72
    rect(s, rx + 0.26, sy, rw - 0.52, 0.46, fill=CYAN, fill_alpha=0.12,
         line=CYAN, lw=1.0, line_alpha=0.35, radius=0.08)
    txt(s, rx + 0.46, sy, rw - 0.9, 0.46, "HUMAN DETECTED", size=15,
        color=CYAN, bold=True, anchor=MSO_ANCHOR.MIDDLE, spacing=1.6)

    # metric rows
    rows = [("Room status", "OCCUPIED", CYAN),
            ("Activity", "WALKING", CYAN),
            ("Alert status", "NORMAL", GREEN)]
    my = sy + 0.62
    for label, value, col in rows:
        txt(s, rx + 0.26, my, 1.70, 0.24, label, size=10, color=DIM)
        txt(s, rx + rw - 2.10, my, 1.84, 0.24, value, size=11.5, color=col,
            bold=True, align=PP_ALIGN.RIGHT, spacing=0.8)
        rect(s, rx + 0.26, my + 0.26, rw - 0.52, 0.008, fill=LINE)
        my += 0.40

    # confidence bar
    txt(s, rx + 0.26, my + 0.02, 2.0, 0.24, "Confidence", size=10, color=DIM)
    txt(s, rx + rw - 2.10, my - 0.02, 1.84, 0.26, "0.87", size=13, color=CYAN,
        bold=True, align=PP_ALIGN.RIGHT)
    bx = rx + 0.26
    bw = rw - 0.52
    rect(s, bx, my + 0.30, bw, 0.14, fill=LINE)
    rect(s, bx, my + 0.30, bw * 0.87, 0.14, fill=CYAN)
    glow(rect(s, bx + bw * 0.87 - 0.05, my + 0.24, 0.10, 0.26, fill=CYAN),
         CYAN, radius_pt=8, alpha=0.8)

    # CSI plot
    gy = my + 0.62
    gh = 0.92
    rect(s, rx + 0.26, gy, rw - 0.52, gh, fill=BG, fill_alpha=0.55,
         line=LINE, lw=0.75, radius=0.06)
    txt(s, rx + 0.40, gy + 0.06, 2.0, 0.20, "CSI  ·  last 6 s", size=8.5,
        color=DIM, spacing=0.8)
    _plot_csi(s, rx + 0.40, gy + 0.30, rw - 0.92, gh - 0.42)

    # links
    ly = gy + gh + 0.18
    txt(s, rx + 0.26, ly, 2.0, 0.20, "LINKS", size=8.5, color=DIM, spacing=1.2)
    for i, (name, strength) in enumerate([("A", 0.82), ("B", 0.64), ("C", 0.31)]):
        lx_i = rx + 0.26 + i * 1.30
        txt(s, lx_i, ly + 0.20, 0.3, 0.20, name, size=9, color=MUTED, bold=True)
        rect(s, lx_i + 0.22, ly + 0.26, 0.90, 0.07, fill=LINE)
        rect(s, lx_i + 0.22, ly + 0.26, 0.90 * strength, 0.07,
             fill=CYAN if strength > 0.5 else AMBER)
    txt(s, rx + rw - 2.10, ly + 0.20, 1.84, 0.20, "No camera feed",
        size=8.5, color=DIM, align=PP_ALIGN.RIGHT, italic=True)

    # bottom note + hardware shot under the steps
    ny = cy + 0.02
    card(s, lx, ny, lw, 0.88, radius=0.10, fill=AMBER, fill_alpha=0.07,
         line=AMBER, line_alpha=0.28)
    picture(s, "s09_hardware.png", lx + 0.03, ny + 0.03, 2.60, 0.82)
    txt(s, lx + 2.86, ny, lw - 3.12, 0.88,
        "Step 5 is the one that matters: the bag changes the signal, and the "
        "system must still say “not human”.",
        size=11.5, color=AMBER, bold=True, anchor=MSO_ANCHOR.MIDDLE,
        line_spacing=1.12)
    footer(s, 9)
    return s


def _plot_csi(slide, x, y, w, h, seed=7):
    """Draw two synthetic CSI traces (quiet → human event → quiet)."""
    rnd = random.Random(seed)
    n = 190

    def series(scale: float, offset: float) -> list[tuple[float, float]]:
        pts = []
        for i in range(n):
            t = i / (n - 1)
            # envelope: quiet, then a walking burst, then quiet again
            env = math.exp(-((t - 0.52) ** 2) / 0.012) * 1.0
            env += math.exp(-((t - 0.30) ** 2) / 0.006) * 0.35
            base = math.sin(t * 26 + offset) * 0.16 * (0.35 + env)
            mid = math.sin(t * 61 + offset * 2) * 0.10 * (0.3 + env)
            noise = rnd.uniform(-0.05, 0.05)
            val = (base + mid + noise) * scale * (0.45 + env)
            pts.append((x + w * t, y + h / 2 - val * (h / 2) * 0.86))
        return pts

    poly(slide, series(1.00, 0.0), color=CYAN, width=1.5)
    poly(slide, series(0.62, 2.1), color=VIOLET, width=1.1, alpha=0.75)
    # event marker
    mx = x + w * 0.52
    rect(slide, mx, y + 0.04, 0.008, h - 0.08, fill=CYAN, fill_alpha=0.35)
    txt(slide, mx - 0.6, y + h - 0.18, 1.2, 0.20, "human event", size=7,
        color=CYAN, align=PP_ALIGN.CENTER)


def slide_impact(prs):
    s = blank(prs)
    backdrop(s, "s10_closing.png")

    txt(s, M, 0.62, 8.0, 0.24, "09  —  IMPACT  &  FUTURE", size=10.5,
        color=CYAN, bold=True, spacing=2.2)
    rect(s, M - 0.22, 0.675, 0.10, 0.10, fill=CYAN)
    txt(s, M, 0.92, CW, 0.52, "Why this matters — and what comes next",
        size=30, color=TEXT, bold=True)

    panel_w = (CW - 0.34) / 2
    py = 1.86
    ph = 3.02

    # ---- impact panel
    card(s, M, py, panel_w, ph, radius=0.14, fill=PANEL, fill_alpha=0.72,
         line=CYAN, line_alpha=0.20)
    rect(s, M, py, panel_w, 0.045, fill=CYAN)
    txt(s, M + 0.34, py + 0.26, panel_w - 0.6, 0.30, "IMPACT", size=12,
        color=CYAN, bold=True, spacing=2.4)
    impacts = [
        ("Camera-free sensing", "Works where cameras are unwanted"),
        ("Privacy by design", "No images, no audio, no identity"),
        ("Low-cost hardware", "ESP32 boards — no specialist RF gear"),
        ("No wearables", "Nothing to wear, charge or remember"),
        ("Edge / near-device processing", "Raw CSI is processed locally"),
        ("Scales room by room", "Add links, reuse the same pipeline"),
    ]
    iy = py + 0.66
    for t, sub in impacts:
        oval(s, M + 0.36, iy + 0.075, 0.10, 0.10, fill=CYAN)
        txt(s, M + 0.60, iy - 0.02, panel_w - 1.0, 0.24, t, size=12.5,
            color=TEXT, bold=True)
        txt(s, M + 0.60, iy + 0.20, panel_w - 1.0, 0.22, sub, size=10,
            color=MUTED)
        iy += 0.38

    # ---- future panel
    fx = M + panel_w + 0.34
    card(s, fx, py, panel_w, ph, radius=0.14, fill=PANEL, fill_alpha=0.72,
         line=VIOLET, line_alpha=0.20)
    rect(s, fx, py, panel_w, 0.045, fill=VIOLET)
    txt(s, fx + 0.34, py + 0.26, panel_w - 0.6, 0.30, "FUTURE", size=12,
        color=VIOLET, bold=True, spacing=2.4)
    futures = [
        ("Reliable fall detection", "Sudden-drop + long-stillness events"),
        ("Multi-room sensing", "Several links, one coordinated model"),
        ("Stronger rejection", "Fans, doors, pets, moving furniture"),
        ("TinyML on the ESP32", "Inference on the board, not the laptop"),
        ("Adaptive calibration", "Self-tuning to each new environment"),
    ]
    fy = py + 0.66
    for t, sub in futures:
        oval(s, fx + 0.36, fy + 0.075, 0.10, 0.10, fill=VIOLET)
        txt(s, fx + 0.60, fy - 0.02, panel_w - 1.0, 0.24, t, size=12.5,
            color=TEXT, bold=True)
        txt(s, fx + 0.60, fy + 0.20, panel_w - 1.0, 0.22, sub, size=10,
            color=MUTED)
        fy += 0.38

    txt(s, fx + 0.34, py + 2.62, panel_w - 0.68, 0.30,
        "CSI sensing is established research. Our contribution is a practical, "
        "low-cost implementation with explicit non-human rejection.",
        size=9.5, color=DIM, italic=True, line_spacing=1.15)

    # ---- closing
    cy = 5.24
    rect(s, M, cy, CW, 0.012, fill=LINE)
    txt(s, M, cy + 0.24, CW - 1.4, 0.60,
        "WiFi is everywhere.  Why not use it to sense the environment — without seeing it?",
        size=20, color=CYAN, bold=True, spacing=0.2)
    txt(s, M, cy + 0.96, CW, 0.30,
        "WiFiSense  ·  Sense movement.  Not identities.",
        size=12.5, color=MUTED, spacing=1.2)
    footer(s, 10)
    return s


# ----------------------------------------------------------- speaker notes

NOTES = {
    1: "Open with the line: “We can tell someone is in this room without ever "
       "taking a picture of them.”\n\n"
       "WiFiSense reads WiFi Channel State Information from two or three ESP32 "
       "boards. No camera, no microphone, nothing to wear. Total hardware cost "
       "is a couple of dev boards and USB cables.\n\n"
       "Say the tagline out loud: Sense movement. Not identities.",
    2: "Four reasons cameras fall short, in order of how much people care: "
       "privacy, blind spots, darkness and glare, and the fact that wearables "
       "get forgotten or refused.\n\n"
       "Land the strip at the bottom: we still need to know whether a space is "
       "occupied — we just should not have to record who is in it.",
    3: "Walk the chain left to right: an ESP32 transmits, the signal crosses the "
       "room, a body disturbs it, a second ESP32 records the change as CSI, and "
       "the model reads the pattern.\n\n"
       "Emphasise “ordinary WiFi frames” — we are not adding any new "
       "transmitting hardware to the room.",
    4: "CSI in one sentence: it is how the signal arrives, not the data it "
       "carries — amplitude and phase per subcarrier, per packet.\n\n"
       "The key idea is step 5: the model reads how those numbers evolve over "
       "seconds, never a single reading.",
    5: "This is the slide that separates us from a naive demo. A bag, a chair, "
       "a fan and a door all change the channel too.\n\n"
       "Call out the line: CSI detects changes in the radio environment, not "
       "humans directly.\n\n"
       "Our three defences: temporal patterns, balanced training data that "
       "includes non-human events, and cross-checking across 2–3 links.",
    6: "Seven small stages — nothing exotic. Keep the vocabulary light: "
       "filter, normalise, window, extract a handful of features, classify.\n\n"
       "Point at the probability bars: the output is a distribution, not a "
       "verdict. Below threshold we say “uncertain”.",
    7: "Left column is radio hardware, right column is software, the arrow is "
       "just numbers moving over a cable.\n\n"
       "Do not skip the green badge: no camera, no microphone, no video or "
       "audio is ever captured or stored. Raw CSI is processed locally.",
    8: "Six places this fits. Keep it to one sentence each — the pictures "
       "carry the slide.\n\n"
       "Read the amber strip honestly. Judges reward teams that know their "
       "limits: results depend on walls, layout and furniture.",
    9: "Ninety seconds, five steps. Steps 1–4 are the happy path; step 5 is "
       "the one to watch: move a bag or chair and the signal still changes, but "
       "the system must not call it a human.\n\n"
       "Dashboard shows room status, activity, confidence score and alert "
       "state — plus per-link signal strength.",
    10: "Impact on the left, roadmap on the right. Say plainly that WiFi sensing "
        "itself is established research — our contribution is a low-cost "
        "implementation with explicit non-human rejection.\n\n"
        "Close on the quote: WiFi is everywhere. Why not use it to sense the "
        "environment — without seeing it?",
}


def add_notes(slide, number: int) -> None:
    text = NOTES.get(number)
    if text:
        slide.notes_slide.notes_text_frame.text = text


def add_transitions(prs) -> None:
    """Subtle fade between every slide (advance stays on click)."""
    for slide in prs.slides:
        sld = slide._element
        anchor = sld.find(qn("p:clrMapOvr"))
        if anchor is None:
            anchor = sld.find(qn("p:cSld"))
        if anchor is None:
            continue
        transition = parse_xml(
            f'<p:transition {nsdecls("p")} spd="fast" advClick="1">'
            f'<p:fade thruBlk="0"/></p:transition>')
        anchor.addnext(transition)


# ------------------------------------------------------------------- main


def main() -> None:
    prs = Presentation()
    prs.slide_width = Inches(SW)
    prs.slide_height = Inches(SH)

    builders = [slide_title, slide_problem, slide_solution, slide_csi,
                slide_false_positive, slide_pipeline, slide_architecture,
                slide_usecases, slide_demo, slide_impact]
    for number, builder in enumerate(builders, start=1):
        add_notes(builder(prs), number)

    add_transitions(prs)

    OUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(OUT))
    size_mb = OUT.stat().st_size / 1e6
    print(f"saved {OUT}  ({len(prs.slides.__iter__.__self__._sldIdLst)} slides, "
          f"{size_mb:.1f} MB)")


if __name__ == "__main__":
    main()
