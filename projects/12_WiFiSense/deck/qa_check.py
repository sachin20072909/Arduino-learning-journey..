"""
qa_check.py — static layout audit for the generated deck.

Renders nothing; it inspects the .pptx geometry and flags:
  * shapes that fall outside the slide
  * pictures whose display aspect differs from the source PNG (distortion)
  * text boxes whose estimated wrapped height exceeds the box height
  * overlapping text boxes (a common cause of "two labels on top of each other")

    python3 deck/qa_check.py
"""

from __future__ import annotations

import sys
from pathlib import Path

from PIL import Image, ImageFont
from pptx import Presentation

ROOT = Path(__file__).resolve().parents[1]
DECK = ROOT / "WiFiSense_Hackathon_Deck.pptx"
EMU_IN = 914400

FONT_PATH = "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"
FONT_BOLD_PATH = "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"
# DejaVu runs ~8% wider than Segoe UI; scale measurements down to compensate.
CALIBRATION = 0.92

_cache: dict[tuple[str, int], ImageFont.FreeTypeFont] = {}


def font(size_pt: int, bold: bool) -> ImageFont.FreeTypeFont:
    key = ("b" if bold else "r", size_pt)
    if key not in _cache:
        px = max(4, round(size_pt * 96 / 72))
        _cache[key] = ImageFont.truetype(
            FONT_BOLD_PATH if bold else FONT_PATH, px)
    return _cache[key]


def text_width(text: str, size_pt: int, bold: bool, spacing_pt: float) -> float:
    """Width of ``text`` in inches, measured in a DejaVu stand-in for Segoe UI."""
    f = font(size_pt, bold)                        # rendered at px = pt * 96/72
    w_px = f.getlength(text) * CALIBRATION
    w_in = w_px / 96.0                             # px -> inches
    if spacing_pt:                                 # letter spacing, in points
        w_in += len(text) * spacing_pt / 72.0
    return w_in


def wrap_lines(text: str, size_pt: int, bold: bool, box_w_in: float,
               spacing_pt: float) -> int:
    lines = 0
    for para in text.split("\n"):
        words = para.split()
        if not words:
            lines += 1
            continue
        cur, count = "", 1
        for word in words:
            trial = f"{cur} {word}".strip()
            if text_width(trial, size_pt, bold, spacing_pt) <= box_w_in:
                cur = trial
            else:
                count += 1
                cur = word
        lines += count
    return lines


def main() -> int:
    prs = Presentation(str(DECK))
    sw, sh = prs.slide_width / EMU_IN, prs.slide_height / EMU_IN
    problems: list[str] = []

    for idx, slide in enumerate(prs.slides, start=1):
        texts: list[tuple[str, float, float, float, float]] = []

        for shp in slide.shapes:
            x, y = shp.left / EMU_IN, shp.top / EMU_IN
            w, h = shp.width / EMU_IN, shp.height / EMU_IN

            # 1. out of bounds
            if x < -0.02 or y < -0.02 or x + w > sw + 0.02 or y + h > sh + 0.02:
                problems.append(
                    f"S{idx:02d} OUT OF BOUNDS  {shp.shape_type} "
                    f"({x:.2f},{y:.2f}) {w:.2f}x{h:.2f}")

            # 2. picture aspect distortion
            if shp.shape_type == 13 or shp.__class__.__name__ == "Picture":
                try:
                    src = Path(shp.image.filename or "")
                    blob = shp.image.blob
                    import io
                    im = Image.open(io.BytesIO(blob))
                    native = im.width / im.height
                    shown = w / h
                    if abs(native - shown) / native > 0.03:
                        problems.append(
                            f"S{idx:02d} IMAGE STRETCH  {src.name or '?'} "
                            f"native {native:.3f} vs shown {shown:.3f}")
                except Exception as exc:                     # noqa: BLE001
                    problems.append(f"S{idx:02d} image error {exc}")

            # 3. text overflow
            if shp.has_text_frame and shp.text_frame.text.strip():
                tf = shp.text_frame
                for p in tf.paragraphs:
                    if not p.runs:
                        continue
                    run = p.runs[0]
                    size = run.font.size.pt if run.font.size else 18
                    bold = bool(run.font.bold)
                    rPr = run._r.find(
                        "{http://schemas.openxmlformats.org/drawingml/2006/main}rPr")
                    spc = 0.0
                    if rPr is not None and rPr.get("spc"):
                        spc = int(rPr.get("spc")) / 100.0
                    ls = p.line_spacing if isinstance(p.line_spacing, float) else 1.2
                    n_lines = wrap_lines(shp.text_frame.text, size, bold, w, spc)
                    needed = n_lines * size * (ls if ls else 1.2) * 1.02 / 72
                    if needed > h + 0.10:
                        problems.append(
                            f"S{idx:02d} TEXT OVERFLOW  +{needed - h:.2f}in  "
                            f"{size:.0f}pt  {n_lines}L  "
                            f"'{shp.text_frame.text[:46].replace(chr(10), ' ')}'")
                    break

                # use the *rendered* width, not the box width, so that
                # intentionally wide boxes do not trigger false overlaps
                first_run = None
                for p in tf.paragraphs:
                    if p.runs:
                        first_run = p.runs[0]
                        break
                if first_run is not None:
                    rsize = first_run.font.size.pt if first_run.font.size else 18
                    rbold = bool(first_run.font.bold)
                    rPr = first_run._r.find(
                        "{http://schemas.openxmlformats.org/"
                        "drawingml/2006/main}rPr")
                    rspc = (int(rPr.get("spc")) / 100.0
                            if (rPr is not None and rPr.get("spc")) else 0.0)
                    longest = max(
                        tf.text.split("\n"),
                        key=lambda ln: text_width(ln, rsize, rbold, 0.0),
                        default="")
                    tw = text_width(longest, rsize, rbold, rspc)
                    # PowerPoint wraps inside the box, so the rendered width is
                    # never wider than the box itself.
                    tw = min(tw, w)
                    align = tf.paragraphs[0].alignment
                    if align == 2:      # center
                        x = x + (w - tw) / 2
                    elif align == 3:    # right
                        x = x + w - tw
                    w = tw

                texts.append((shp.text_frame.text[:30], x, y, w, h))

        # 4. overlapping text boxes
        for i in range(len(texts)):
            for j in range(i + 1, len(texts)):
                _, ax, ay, aw, ah = texts[i]
                _, bx, by, bw, bh = texts[j]
                ox = max(0.0, min(ax + aw, bx + bw) - max(ax, bx))
                oy = max(0.0, min(ay + ah, by + bh) - max(ay, by))
                if ox > 0.12 and oy > 0.10:
                    problems.append(
                        f"S{idx:02d} TEXT OVERLAP  {ox:.2f}x{oy:.2f}in  "
                        f"'{texts[i][0]}' <> '{texts[j][0]}'")

    if problems:
        print(f"{len(problems)} potential issue(s):\n")
        for p in problems:
            print("  " + p)
    else:
        print("clean — no geometry problems detected")
    return len(problems)


if __name__ == "__main__":
    sys.exit(0 if main() == 0 else 1)
