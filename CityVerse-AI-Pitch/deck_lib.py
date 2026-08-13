"""
CityVerse AI — design system for the hackathon deck.
Theme helpers, glassmorphism, glows, gradients, and real PowerPoint
animation/transition XML injection (Morph, Fade, Wipe, Zoom, motion paths,
looping pulses) on top of python-pptx.
"""
from lxml import etree
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.oxml.ns import qn

# ------------------------------------------------------------------ palette
BG      = "050A18"   # near-black navy
BG_SOFT = "0A1226"
CARD    = "0D1B33"
CYAN    = "22D3EE"
BLUE    = "3B82F6"
PURPLE  = "8B5CF6"
ICE     = "9FD8FF"
WHITE   = "F1F5F9"
MUTED   = "8EA3C0"
DIM     = "5A6E8C"
RED     = "F43F5E"
AMBER   = "F59E0B"
GREEN   = "34D399"
GRIDLN  = "1B2B4A"

P_NS  = "http://schemas.openxmlformats.org/presentationml/2006/main"
A_NS  = "http://schemas.openxmlformats.org/drawingml/2006/main"
R_NS  = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
MC_NS = "http://schemas.openxmlformats.org/markup-compatibility/2006"
P14_NS = "http://schemas.microsoft.com/office/powerpoint/2010/main"

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)

def _sub(parent, tag, **attrs):
    el = etree.SubElement(parent, tag)
    for k, v in attrs.items():
        el.set(k, str(v))
    return el

def _xfrag(xml_inner, root="p:x"):
    """Parse an XML fragment; returns the fragment's first (real) element."""
    wrapped = ('<{0} xmlns:p="{1}" xmlns:a="{2}" xmlns:mc="{3}" xmlns:p14="{4}">'
               '{5}</{0}>').format(root, P_NS, A_NS, MC_NS, P14_NS, xml_inner)
    return etree.fromstring(wrapped)[0]

# ------------------------------------------------------------------ fills / lines
def _srgb(color, alpha=None):
    s = '<a:srgbClr val="%s">' % color
    if alpha is not None:
        s += '<a:alpha val="%d"/>' % int(alpha * 1000)
    return s + '</a:srgbClr>'

def fill_solid(shape, color, alpha=None):
    spPr = shape._element.spPr
    for tag in ("a:noFill", "a:solidFill", "a:gradFill", "a:blipFill",
                "a:pattFill", "a:grpFill"):
        for el in spPr.findall(qn(tag)):
            spPr.remove(el)
    node = _xfrag('<a:solidFill>%s</a:solidFill>' % _srgb(color, alpha))
    ln = spPr.find(qn("a:ln"))
    if ln is not None:
        ln.addprevious(node)
    else:
        spPr.append(node)

def fill_gradient(shape, stops, angle_deg=90.0):
    """stops: [(pos 0-100, color, alpha 0-100 or None)], angle clockwise from 3 o'clock."""
    spPr = shape._element.spPr
    for tag in ("a:noFill", "a:solidFill", "a:gradFill", "a:blipFill",
                "a:pattFill", "a:grpFill"):
        for el in spPr.findall(qn(tag)):
            spPr.remove(el)
    gs = "".join('<a:gs pos="%d">%s</a:gs>' % (int(p * 1000), _srgb(c, a))
                 for p, c, a in stops)
    node = _xfrag('<a:gradFill><a:gsLst>%s</a:gsLst>'
                  '<a:lin ang="%d" scaled="1"/></a:gradFill>'
                  % (gs, int(angle_deg * 60000)))
    ln = spPr.find(qn("a:ln"))
    if ln is not None:
        ln.addprevious(node)
    else:
        spPr.append(node)

def fill_radial(shape, stops):
    """Radial gradient (center -> edge), used for ambient glow blobs."""
    spPr = shape._element.spPr
    for tag in ("a:noFill", "a:solidFill", "a:gradFill", "a:blipFill",
                "a:pattFill", "a:grpFill"):
        for el in spPr.findall(qn(tag)):
            spPr.remove(el)
    gs = "".join('<a:gs pos="%d">%s</a:gs>' % (int(p * 1000), _srgb(c, a))
                 for p, c, a in stops)
    node = _xfrag('<a:gradFill><a:gsLst>%s</a:gsLst>'
                  '<a:path path="circle"><a:fillToRect l="50000" t="50000" '
                  'r="50000" b="50000"/></a:path></a:gradFill>' % gs)
    ln = spPr.find(qn("a:ln"))
    if ln is not None:
        ln.addprevious(node)
    else:
        spPr.append(node)

def no_line(shape):
    ln = shape._element.spPr.get_or_add_ln()
    for el in list(ln):
        ln.remove(el)
    ln.append(_xfrag('<a:noFill/>'))

def line_style(shape, color, width_pt=1.0, alpha=None, dash=None,
               cap="rnd"):
    ln = shape._element.spPr.get_or_add_ln()
    for el in list(ln):
        ln.remove(el)
    ln.set("w", str(int(width_pt * 12700)))
    ln.set("cap", cap)
    ln.append(_xfrag('<a:solidFill>%s</a:solidFill>' % _srgb(color, alpha)))
    if dash:
        ln.append(_xfrag('<a:prstDash val="%s"/>' % dash))

def no_shadow(shape):
    try:
        shape.shadow.inherit = False
    except Exception:
        pass

def add_glow(shape, color, radius_pt=10, alpha=60):
    spPr = shape._element.spPr
    for el in spPr.findall(qn("a:effectLst")):
        spPr.remove(el)
    spPr.append(_xfrag('<a:effectLst><a:glow rad="%d">%s</a:glow></a:effectLst>'
                       % (int(radius_pt * 12700), _srgb(color, alpha))))
    no_shadow(shape)

def add_shadow(shape, blur=18, dist=8, alpha=45, color="000000"):
    spPr = shape._element.spPr
    for el in spPr.findall(qn("a:effectLst")):
        spPr.remove(el)
    spPr.append(_xfrag(
        '<a:effectLst><a:outerShdw blurRad="%d" dist="%d" dir="5400000" '
        'rotWithShape="0">%s</a:outerShdw></a:effectLst>'
        % (blur * 12700, dist * 12700, _srgb(color, alpha))))
    no_shadow(shape)

# ------------------------------------------------------------------ shapes
def rect(slide, x, y, w, h, color=None, alpha=None, line=None,
         line_w=1.0, line_alpha=None, radius=None, glow=None,
         glow_r=10, dash=None, shape_type=MSO_SHAPE.RECTANGLE,
         grad=None, radial=None, shadow=False, rot=None, name=None):
    if radius is not None:
        shape_type = MSO_SHAPE.ROUNDED_RECTANGLE
    sp = slide.shapes.add_shape(shape_type, Inches(x), Inches(y),
                                Inches(w), Inches(h))
    no_shadow(sp)
    if radius is not None:
        try:
            sp.adjustments[0] = radius
        except Exception:
            pass
    if grad:
        fill_gradient(sp, grad[0], grad[1])
    elif radial:
        fill_radial(sp, radial)
    elif color:
        fill_solid(sp, color, alpha)
    else:
        sp.fill.background()
    if line:
        line_style(sp, line, line_w, line_alpha, dash)
    else:
        no_line(sp)
    if glow:
        add_glow(sp, glow, glow_r)
    if shadow:
        add_shadow(sp)
    if rot is not None:
        sp.rotation = rot
    if name:
        sp.name = name
    sp.text_frame.paragraphs[0].text = ""
    return sp

def glass(slide, x, y, w, h, radius=0.16, tint=CARD, alpha=55,
          edge=CYAN, edge_alpha=22, edge_w=1.0, glow=None, glow_r=8,
          name=None):
    """Signature glassmorphism card."""
    return rect(slide, x, y, w, h, color=tint, alpha=alpha, radius=radius,
                line=edge, line_alpha=edge_alpha, line_w=edge_w,
                glow=glow, glow_r=glow_r, name=name)

def oval(slide, cx, cy, d, color=None, alpha=None, line=None, line_w=1.0,
         line_alpha=None, glow=None, glow_r=8, name=None):
    sp = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(cx - d / 2),
                                Inches(cy - d / 2), Inches(d), Inches(d))
    no_shadow(sp)
    if color:
        fill_solid(sp, color, alpha)
    else:
        sp.fill.background()
    if line:
        line_style(sp, line, line_w, line_alpha)
    else:
        no_line(sp)
    if glow:
        add_glow(sp, glow, glow_r)
    if name:
        sp.name = name
    return sp

def blob(slide, cx, cy, d, color, alpha_center=30):
    """Soft radial light."""
    sp = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(cx - d / 2),
                                Inches(cy - d / 2), Inches(d), Inches(d))
    no_shadow(sp)
    no_line(sp)
    fill_radial(sp, [(0, color, alpha_center), (65, color, alpha_center // 3),
                     (100, color, 0)])
    return sp

def seg(slide, x1, y1, x2, y2, color=CYAN, w=1.0, alpha=40, dash=None,
        arrow=False, glow=None, glow_r=6):
    ln = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1),
                                    Inches(y1), Inches(x2), Inches(y2))
    no_shadow(ln)
    line_style(ln, color, w, alpha, dash)
    lnEl = ln._element.spPr.find(qn("a:ln"))
    if arrow:
        lnEl.append(_xfrag('<a:tailEnd type="triangle" w="med" len="med"/>'))
    if glow:
        add_glow(ln, glow, glow_r)
    ln.shadow.inherit = False
    return ln

def pic(slide, path, x, y, w, h, radius=None, line=None, line_w=1.25,
        line_alpha=45, glow=None, glow_r=12, name=None):
    p = slide.shapes.add_picture(path, Inches(x), Inches(y), Inches(w),
                                 Inches(h))
    p.shadow.inherit = False
    if radius is not None:
        spPr = p._element.spPr
        for el in spPr.findall(qn("a:prstGeom")):
            spPr.remove(el)
        spPr.append(_xfrag('<a:prstGeom prst="roundRect"><a:avLst>'
                           '<a:gd name="adj" fmla="val %d"/></a:avLst>'
                           '</a:prstGeom>' % int(radius * 100000)))
    if line:
        line_style(p, line, line_w, line_alpha)
    if glow:
        add_glow(p, glow, glow_r)
    if name:
        p.name = name
    return p

# ------------------------------------------------------------------ text
FONT      = "Segoe UI"
FONT_LT   = "Segoe UI Light"
FONT_SB   = "Segoe UI Semibold"

def _apply_run(r, text, size, color, bold=False, italic=False, font=FONT,
               spc=None, alpha=None, grad=None):
    r.text = text
    f = r.font
    f.size = Pt(size)
    f.bold = bold
    f.italic = italic
    f.name = font
    f.color.rgb = RGBColor.from_string(color)
    rPr = r._r.get_or_add_rPr()
    if spc:
        rPr.set("spc", str(int(spc)))
    if alpha is not None:
        sf = rPr.find(qn("a:solidFill"))
        if sf is not None:
            clr = sf.find(qn("a:srgbClr"))
            if clr is not None:
                clr.append(_xfrag('<a:alpha val="%d"/>' % int(alpha * 1000)))
    if grad:
        sf = rPr.find(qn("a:solidFill"))
        if sf is not None:
            rPr.remove(sf)
        gs = "".join('<a:gs pos="%d">%s</a:gs>' % (int(p * 1000), _srgb(c, None))
                     for p, c in enumerate([]))
        stops, ang = grad
        gs = "".join('<a:gs pos="%d">%s</a:gs>' % (int(p * 1000), _srgb(c, None))
                     for p, c in stops)
        node = _xfrag('<a:gradFill><a:gsLst>%s</a:gsLst>'
                      '<a:lin ang="%d" scaled="1"/></a:gradFill>'
                      % (gs, int(ang * 60000)))
        anchor = None
        for t in ("a:latin", "a:ea", "a:cs", "a:sym", "a:hlinkClick",
                  "a:hlinkMouseOver"):
            anchor = rPr.find(qn(t))
            if anchor is not None:
                break
        if anchor is not None:
            anchor.addprevious(node)
        else:
            rPr.append(node)

ALIGN = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT}
ANCHOR = {"t": MSO_ANCHOR.TOP, "m": MSO_ANCHOR.MIDDLE, "b": MSO_ANCHOR.BOTTOM}

def txt(slide, x, y, w, h, runs, align="l", anchor="t", wrap=True,
        line_spacing=None, name=None, shrink=False):
    """runs: list of paragraphs; each paragraph is a list of run dicts."""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = wrap
    tf.vertical_anchor = ANCHOR[anchor]
    for m in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
        setattr(tf, m, 0)
    first = True
    for para in runs:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.alignment = ALIGN[align]
        if line_spacing:
            p.line_spacing = line_spacing
        if isinstance(para, dict):
            if "space_before" in para: p.space_before = Pt(para["space_before"])
            if "space_after" in para:  p.space_after  = Pt(para["space_after"])
            if "align" in para:        p.alignment = ALIGN[para["align"]]
            para_runs = para["runs"]
        else:
            para_runs = para
        for rd in para_runs:
            _apply_run(p.add_run(), rd.get("t", ""), rd.get("s", 12),
                       rd.get("c", WHITE), rd.get("b", False),
                       rd.get("i", False), rd.get("f", FONT),
                       rd.get("spc"), rd.get("a"), rd.get("g"))
    if name:
        tb.name = name
    return tb

def R(t, s=12, c=WHITE, b=False, i=False, f=FONT, spc=None, a=None, g=None):
    return {"t": t, "s": s, "c": c, "b": b, "i": i, "f": f, "spc": spc,
            "a": a, "g": g}

# ------------------------------------------------------------------ animations
class Anim:
    """Collects effects for one slide and renders <p:timing>."""

    def __init__(self):
        self.effects = []      # list of dicts
        self._id = [5]

    def _next(self):
        self._id[0] += 1
        return self._id[0]

    def _spid(self, shape):
        return shape._element.get(  # nvSpPr/cNvPr or nvPicPr/cNvPr share 'id'
            "id") if shape._element.tag.endswith("}sp") else None

    def add(self, shape, kind, dur=600, delay=0, start="after",
            direction=None, path=None, loop=False, fade_dur=None):
        spid = shape.shape_id
        node = {"click": "clickEffect",
                "with": "withEffect",
                "after": "afterEffect"}[start]
        if kind == "rise":
            # paired fade-in + upward motion
            self.add(shape, "motion", dur=dur, delay=delay, start=start,
                     path=path or "M 0.0 0.045 L 0.0 0.0")
            self.effects.append(dict(kind="fade", spid=spid, dur=dur,
                                     delay=delay, node="withEffect",
                                     loop=False, direction=None, path=None))
            return
        if kind == "flow":      # looping motion along a path
            self.effects.append(dict(kind="motion", spid=spid, dur=dur,
                                     delay=delay, node=node, loop=True,
                                     direction=None, path=path))
            self.effects.append(dict(kind="breathe", spid=spid, dur=dur,
                                     delay=delay, node="withEffect",
                                     loop=True, direction=None, path=None))
            return
        self.effects.append(dict(kind=kind, spid=spid, dur=dur, delay=delay,
                                 node=node, loop=loop, direction=direction,
                                 path=path))

    # ---------------------------------------------------------- xml pieces
    def _set_visible(self, spid):
        nid = self._next()
        return ('<p:set><p:cBhvr><p:cTn id="%d" dur="1" fill="hold">'
                '<p:stCondLst><p:cond delay="0"/></p:stCondLst></p:cTn>'
                '<p:tgtEl><p:spTgt spid="%d"/></p:tgtEl>'
                '<p:attrNameLst><p:attrName>style.visibility</p:attrName>'
                '</p:attrNameLst></p:cBhvr>'
                '<p:to><p:strVal val="visible"/></p:to></p:set>') % (nid, spid)

    def _anim_effect(self, spid, transition, flt, dur, loop=False):
        nid = self._next()
        rep = ' repeatCount="indefinite"' if loop else ""
        return ('<p:animEffect transition="%s" filter="%s"><p:cBhvr>'
                '<p:cTn id="%d" dur="%d"%s/><p:tgtEl><p:spTgt spid="%d"/>'
                '</p:tgtEl></p:cBhvr></p:animEffect>'
                % (transition, flt, nid, dur, rep, spid))

    def _anim_motion(self, spid, path, dur, loop=False):
        nid = self._next()
        rep = ' repeatCount="indefinite"' if loop else ""
        return ('<p:animMotion origin="layout" path="%s" '
                'pathEditMode="relative" rAng="0" ptsTypes=""><p:cBhvr>'
                '<p:cTn id="%d" dur="%d" fill="hold"%s/>'
                '<p:tgtEl><p:spTgt spid="%d"/></p:tgtEl></p:cBhvr>'
                '</p:animMotion>' % (path, nid, dur, rep, spid))

    def _anim_scale(self, spid, pct, dur, loop=True):
        nid = self._next()
        rep = ' repeatCount="indefinite"' if loop else ""
        return ('<p:animScale><p:cBhvr><p:cTn id="%d" dur="%d" autoRev="1"%s/>'
                '<p:tgtEl><p:spTgt spid="%d"/></p:tgtEl></p:cBhvr>'
                '<p:by x="%d%%" y="%d%%"/></p:animScale>'
                % (nid, dur, rep, spid, pct, pct))

    def _anim_breathe(self, spid, dur, loop=True):
        nid = self._next()
        rep = ' repeatCount="indefinite"' if loop else ""
        return ('<p:animEffect transition="out" filter="fade"><p:cBhvr>'
                '<p:cTn id="%d" dur="%d" autoRev="1"%s/>'
                '<p:tgtEl><p:spTgt spid="%d"/></p:tgtEl></p:cBhvr>'
                '</p:animEffect>' % (nid, dur, rep, spid))

    def _effect_par(self, eff):
        kind, spid = eff["kind"], eff["spid"]
        dur, delay, node = eff["dur"], eff["delay"], eff["node"]
        preset, children = 'presetID="10" presetClass="entr" presetSubtype="0"', []
        if kind == "fade":
            children = [self._set_visible(spid),
                        self._anim_effect(spid, "in", "fade", dur)]
        elif kind == "zoom":
            preset = 'presetID="53" presetClass="entr" presetSubtype="0"'
            nid = self._next()
            scale = ('<p:animScale><p:cBhvr><p:cTn id="%d" dur="%d"/>'
                     '<p:tgtEl><p:spTgt spid="%d"/></p:tgtEl></p:cBhvr>'
                     '<p:from x="60%%" y="60%%"/><p:to x="100%%" y="100%%"/>'
                     '</p:animScale>' % (nid, dur, spid))
            children = [self._set_visible(spid),
                        self._anim_effect(spid, "in", "fade", dur),
                        scale]
        elif kind == "wipe":
            sub = {"right": ("fromLeft", 2), "left": ("fromRight", 1),
                   "up": ("fromBottom", 8), "down": ("fromTop", 4),
                   None: ("fromBottom", 8)}[eff["direction"]]
            preset = ('presetID="22" presetClass="entr" presetSubtype="%d"'
                      % sub[1])
            children = [self._set_visible(spid),
                        self._anim_effect(spid, "in", "wipe(%s)" % sub[0], dur)]
        elif kind == "wheel":
            preset = 'presetID="21" presetClass="entr" presetSubtype="1"'
            children = [self._set_visible(spid),
                        self._anim_effect(spid, "in", "wheel(1)", dur)]
        elif kind == "motion":
            preset = 'presetID="64" presetClass="path" presetSubtype="0"'
            children = [self._anim_motion(spid, eff["path"], dur, eff["loop"])]
        elif kind == "pulse":
            preset = 'presetID="6" presetClass="emph" presetSubtype="0"'
            children = [self._anim_scale(spid, 116, dur, eff["loop"])]
        elif kind == "blink":
            preset = 'presetID="8" presetClass="emph" presetSubtype="0"'
            children = [self._anim_breathe(spid, dur, eff["loop"])]
        elif kind == "breathe":
            preset = 'presetID="8" presetClass="emph" presetSubtype="0"'
            children = [self._anim_breathe(spid, dur, eff["loop"])]
        elif kind == "fadeout":
            preset = 'presetID="10" presetClass="exit" presetSubtype="0"'
            children = [self._anim_effect(spid, "out", "fade", dur)]
        nid = self._next()
        return ('<p:par><p:cTn id="%d" %s fill="hold" grpId="0" nodeType="%s">'
                '<p:stCondLst><p:cond delay="%d"/></p:stCondLst>'
                '<p:childTnLst>%s</p:childTnLst></p:cTn></p:par>'
                % (nid, preset, node, delay, "".join(children)))

    def build(self):
        if not self.effects:
            return None
        body = "".join(self._effect_par(e) for e in self.effects)
        xml = ('<p:timing><p:tnLst><p:par>'
               '<p:cTn id="1" dur="indefinite" restart="never" nodeType="tmRoot">'
               '<p:childTnLst><p:seq concurrent="1" nextAc="seek">'
               '<p:cTn id="2" dur="indefinite" nodeType="mainSeq">'
               '<p:childTnLst><p:par>'
               '<p:cTn id="3" fill="hold">'
               '<p:stCondLst><p:cond delay="indefinite"/></p:stCondLst>'
               '<p:childTnLst><p:par>'
               '<p:cTn id="4" fill="hold">'
               '<p:stCondLst><p:cond delay="0"/></p:stCondLst>'
               '<p:childTnLst>%s</p:childTnLst>'
               '</p:cTn></p:par></p:childTnLst></p:cTn></p:par>'
               '</p:childTnLst></p:cTn>'
               '<p:prevCondLst><p:cond evt="onPrev" delay="0">'
               '<p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:prevCondLst>'
               '<p:nextCondLst><p:cond evt="onNext" delay="0">'
               '<p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:nextCondLst>'
               '</p:seq></p:childTnLst></p:cTn></p:par></p:tnLst>'
               '<p:bldLst/></p:timing>' % body)
        return _xfrag(xml, root="p:timingX")[0] if False else \
            etree.fromstring(('<p:timing xmlns:p="%s" xmlns:a="%s">%s</p:timing>'
                              % (P_NS, A_NS, xml[len("<p:timing>"):-len("</p:timing>")])))

# ------------------------------------------------------------------ transitions
def add_transition(slide, morph_dur=1100):
    xml = ('<mc:AlternateContent xmlns:mc="%s">'
           '<mc:Choice xmlns:p14="%s" Requires="p14">'
           '<p:transition xmlns:p="%s" spd="slow" p14:dur="%d">'
           '<p14:morph option="byObject"/></p:transition></mc:Choice>'
           '<mc:Fallback><p:transition xmlns:p="%s" spd="slow" advClick="1">'
           '<p:fade/></p:transition></mc:Fallback></mc:AlternateContent>'
           % (MC_NS, P14_NS, P_NS, morph_dur, P_NS))
    node = etree.fromstring(xml)
    sld = slide._element
    clr = sld.find(qn("p:clrMapOvr"))
    if clr is not None:
        clr.addnext(node)
    else:
        sld.find(qn("p:cSld")).addnext(node)

def add_timing(slide, anim):
    node = anim.build()
    if node is None:
        return
    sld = slide._element
    ext = sld.find(qn("p:extLst"))
    if ext is not None:
        ext.addprevious(node)
    else:
        sld.append(node)

# ------------------------------------------------------------------ notes
def add_notes(slide, text):
    nf = slide.notes_slide.notes_text_frame
    nf.text = text

# ------------------------------------------------------------------ furniture
def content_base(slide, kicker, title, idx, kicker_color=CYAN,
                 title_size=33):
    rect(slide, 0, 0, 13.333, 7.5, color=BG, name="bg")
    blob(slide, 12.4, 0.4, 5.2, BLUE, 20)
    blob(slide, 0.7, 7.2, 4.6, PURPLE, 16)
    rect(slide, 0.62, 0.66, 0.30, 0.052, color=kicker_color,
         glow=kicker_color, glow_r=6, name="kbar")
    kb = txt(slide, 1.02, 0.535, 8.5, 0.32,
             [[R(kicker.upper(), 10.5, kicker_color, True, spc=320)]],
             name="kicker")
    tt = txt(slide, 0.62, 0.95, 12.1, 0.72,
             [[R(title, title_size, WHITE, True)]], name="title")
    pg = txt(slide, 11.85, 7.13, 0.86, 0.24,
             [[R("%02d" % idx, 9, MUTED, True),
               R(" / 12", 9, DIM)]], align="r", name="pageno")
    ft = txt(slide, 0.62, 7.13, 6.0, 0.24,
             [[R("CITYVERSE AI", 8, DIM, True, spc=260),
               R("  ·  SMART DIGITAL TWIN PLATFORM", 8, DIM, spc=180)]],
             name="footer")
    return {"kicker": kb, "title": tt, "page": pg, "footer": ft}

def freeform_polyline(slide, pts, color=CYAN, w=2.0, glow=None, closed=False,
                      fill_grad=None):
    """pts in inches [(x,y),...] -> freeform shape."""
    from pptx.util import Emu as _E
    fb = slide.shapes.build_freeform(_E(Inches(pts[0][0])),
                                     _E(Inches(pts[0][1])), scale=1.0)
    fb.add_line_segments([(_E(Inches(px)), _E(Inches(py)))
                          for px, py in pts[1:]], close=closed)
    sp = fb.convert_to_shape()
    no_shadow(sp)
    if fill_grad:
        fill_gradient(sp, fill_grad[0], fill_grad[1])
        no_line(sp)
    else:
        sp.fill.background()
        line_style(sp, color, w)
        if glow:
            add_glow(sp, glow, 8)
    return sp
